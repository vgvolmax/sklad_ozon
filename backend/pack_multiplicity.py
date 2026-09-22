"""Persistent article-level pack multiplicity master data."""

import hashlib
import json
from dataclasses import dataclass, replace
from datetime import datetime, timedelta, timezone
from io import BytesIO
from math import isfinite

from openpyxl import Workbook, load_workbook

from backend.ingestion.normalization import normalize_header
from backend.domain.contracts import ReportMeta
from backend.ingestion.supplier_packaging import (PackMultiplicityEvidence,
                                                   import_supplier_packaging,
                                                   normalize_supplier_article)
from backend.project import PackMultiplicityRecord, Project


@dataclass(frozen=True, slots=True)
class ResolvedPackMultiplicity:
    pack_multiple: int | None
    source: str
    unitka_pack_multiple: int | None
    rtp_price_pack_multiple: int | None
    override_pack_multiple: int | None
    updated_at: str | None


@dataclass(frozen=True, slots=True)
class PackImportDiagnostic:
    row: int
    article: str | None
    code: str
    message: str


@dataclass(frozen=True, slots=True)
class ParsedPackImport:
    source_format: str
    values: dict[str, int]
    diagnostics: tuple[PackImportDiagnostic, ...]


@dataclass(frozen=True, slots=True)
class EffectivePackEvidence:
    """One immutable effective pack value used by an analysis run."""
    article: str
    pack_multiple: int | None
    source: str
    reason_codes: tuple[str, ...]


def validate_pack_multiple(value: object, *, excel: bool = False) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError("Кратность должна быть положительным целым числом.")
    if isinstance(value, float) and (not excel or not isfinite(value) or not value.is_integer()):
        raise ValueError("Кратность должна быть положительным целым числом.")
    result = int(value)
    if result <= 0:
        raise ValueError("Кратность должна быть положительным целым числом.")
    return result


def resolve_pack_multiplicity(record: PackMultiplicityRecord | None) -> ResolvedPackMultiplicity:
    record = record or PackMultiplicityRecord()
    if record.override_pack_multiple is not None:
        return ResolvedPackMultiplicity(record.override_pack_multiple, record.override_origin or "unknown", record.unitka_pack_multiple, record.rtp_price_pack_multiple, record.override_pack_multiple, record.override_updated_at)
    if record.rtp_price_pack_multiple is not None:
        return ResolvedPackMultiplicity(record.rtp_price_pack_multiple, "rtp_price", record.unitka_pack_multiple, record.rtp_price_pack_multiple, None, record.rtp_price_updated_at)
    if record.unitka_pack_multiple is not None:
        return ResolvedPackMultiplicity(record.unitka_pack_multiple, "unitka", record.unitka_pack_multiple, None, None, None)
    return ResolvedPackMultiplicity(None, "unknown", None, None, None, None)


def moscow_now() -> str:
    return datetime.now(timezone(timedelta(hours=3))).isoformat()


def set_override(project: Project, article_value: object, value: object, origin: str, *, updated_at: str | None = None) -> tuple[Project, str]:
    article = normalize_supplier_article(article_value)
    multiple = validate_pack_multiple(value)
    if origin not in {"manual", "import"}: raise ValueError("Unknown override origin.")
    old = project.pack_multiplicity.get(article, PackMultiplicityRecord())
    record = replace(old, override_pack_multiple=multiple, override_origin=origin, override_updated_at=updated_at or moscow_now())
    return replace(project, pack_multiplicity={**project.pack_multiplicity, article: record}), article


def reset_override(project: Project, article_value: object) -> tuple[Project, str]:
    article = normalize_supplier_article(article_value)
    old = project.pack_multiplicity.get(article, PackMultiplicityRecord())
    record = replace(old, override_pack_multiple=None, override_origin=None, override_updated_at=None)
    return replace(project, pack_multiplicity={**project.pack_multiplicity, article: record}), article


def sync_unitka_baseline(project: Project, evidence) -> Project:
    records = dict(project.pack_multiplicity)
    for item in evidence:
        old = records.get(item.article, PackMultiplicityRecord())
        records[item.article] = replace(old, unitka_pack_multiple=item.pack_multiple)
    return replace(project, pack_multiplicity=records)


def pack_multiplicity_fingerprint(project: Project) -> str:
    """Return a deterministic revision for pack master data only."""
    payload = [
        {
            "article": article,
            "unitka_pack_multiple": record.unitka_pack_multiple,
            "rtp_price_pack_multiple": record.rtp_price_pack_multiple,
            "override_pack_multiple": record.override_pack_multiple,
            "override_origin": record.override_origin,
        }
        for article, record in sorted(project.pack_multiplicity.items())
    ]
    canonical = json.dumps(
        payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(canonical).hexdigest()


def build_effective_pack_evidence(project: Project, current_unitka_evidence) -> tuple[EffectivePackEvidence, ...]:
    """Resolve override > RTP price > current valid Unitka > persisted Unitka.

    Current invalid/conflicting Unitka evidence remains causal only when it is
    not masked by a valid override or RTP-price value.
    """
    current = {item.article: item for item in current_unitka_evidence}
    articles = sorted(set(project.pack_multiplicity) | set(current))
    result = []
    for article in articles:
        record = project.pack_multiplicity.get(article)
        resolved = resolve_pack_multiplicity(record)
        observed = current.get(article)
        if resolved.source in {"manual", "import", "rtp_price"}:
            result.append(EffectivePackEvidence(
                article, resolved.pack_multiple, resolved.source, ()))
        elif observed is not None:
            if observed.pack_multiple is not None:
                result.append(EffectivePackEvidence(
                    article, observed.pack_multiple, "unitka", ()))
            else:
                result.append(EffectivePackEvidence(
                    article, None, "unknown",
                    observed.reason_codes or ("MISSING_PACK_MULTIPLICITY",)))
        elif resolved.pack_multiple is not None:
            result.append(EffectivePackEvidence(
                article, resolved.pack_multiple, "unitka", ()))
        else:
            result.append(EffectivePackEvidence(
                article, None, "unknown", ("MISSING_PACK_MULTIPLICITY",)))
    return tuple(result)


def _rtp_import_diagnostic(
        item: PackMultiplicityEvidence) -> PackImportDiagnostic:
    code = item.reason_codes[0]
    if code == "INVALID_PACK_MULTIPLICITY":
        message = (
            f"Кратность коробки {item.source_value!r} не является точным количеством."
            if item.source_value is not None
            else "Кратность коробки не является точным количеством."
        )
    else:
        message = "Для артикула указаны разные кратности внешней коробки."
    return PackImportDiagnostic(
        item.source_row or 0, item.article, code, message)


def parse_import_xlsx(data: bytes) -> ParsedPackImport:
    try: workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc: raise ValueError("Файл должен быть корректным XLSX.") from exc
    try:
        if "Прайс списком" in workbook.sheetnames:
            sheet = workbook["Прайс списком"]
            headers_found = any(
                "код" in (headers := [normalize_header(v) for v in row]) and
                "упак" in headers
                for row in sheet.iter_rows(max_row=30, values_only=True)
            )
            if headers_found:
                result = import_supplier_packaging(
                    data, ReportMeta("rtp-price.xlsx", moscow_now()),
                    workbook=workbook)
                values = {item.article: item.pack_multiple for item in result.records
                          if item.pack_multiple is not None}
                diagnostics = [
                    _rtp_import_diagnostic(item)
                    for item in result.records if item.reason_codes
                ]
                diagnostics.extend(PackImportDiagnostic(
                    diagnostic.row or 0, None, diagnostic.code,
                    diagnostic.message)
                    for diagnostic in result.diagnostics
                    if diagnostic.code == "INVALID_SUPPLIER_ARTICLE")
                return ParsedPackImport("rtp_price", values, tuple(diagnostics))
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        normalized = [normalize_header(value) for value in (header or ())]
        if "артикул" not in normalized or "кратность" not in normalized:
            raise ValueError("Нужны колонки «Артикул» и «Кратность».")
        ai, pi = normalized.index("артикул"), normalized.index("кратность")
        candidates, errors = {}, []
        occurrences: dict[str, list[int]] = {}
        for row_number, row in enumerate(rows, 2):
            av = row[ai] if ai < len(row) else None; pv = row[pi] if pi < len(row) else None
            if av is None and pv is None: continue
            try: article = normalize_supplier_article(av)
            except ValueError:
                errors.append(PackImportDiagnostic(row_number, None if av is None else str(av), "INVALID_ARTICLE", "Артикул должен быть заполнен и корректен.")); continue
            occurrences.setdefault(article, []).append(row_number)
            try: candidates[article] = validate_pack_multiple(pv, excel=True)
            except ValueError as exc: errors.append(PackImportDiagnostic(row_number, article, "INVALID_PACK_MULTIPLICITY", str(exc)))
        duplicates = {a for a, rows_ in occurrences.items() if len(rows_) > 1}
        for article in sorted(duplicates):
            candidates.pop(article, None)
            errors.append(PackImportDiagnostic(occurrences[article][0], article, "DUPLICATE_ARTICLE", "Артикул встречается в файле несколько раз и не импортирован."))
        invalid_articles = {e.article for e in errors if e.article}
        for article in invalid_articles: candidates.pop(article, None)
        return ParsedPackImport("canonical", candidates, tuple(errors))
    finally: workbook.close()


def apply_rtp_price_snapshot(project: Project, values: dict[str, int], *,
                             updated_at: str | None = None) -> Project:
    """Replace only the authoritative RTP-price layer, preserving all others."""
    records = dict(project.pack_multiplicity)
    semantic_current = {article: record.rtp_price_pack_multiple
                        for article, record in records.items()
                        if record.rtp_price_pack_multiple is not None}
    if semantic_current == values:
        return project
    stamp = updated_at or moscow_now()
    for article, record in tuple(records.items()):
        if record.rtp_price_pack_multiple is not None:
            records[article] = replace(record, rtp_price_pack_multiple=None,
                                       rtp_price_updated_at=None)
    for article, multiple in values.items():
        old = records.get(article, PackMultiplicityRecord())
        records[article] = replace(old, rtp_price_pack_multiple=multiple,
                                   rtp_price_updated_at=stamp)
    return replace(project, pack_multiplicity=records)


def export_xlsx(items: list[dict[str, object]]) -> bytes:
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Кратность упаковки"
    sheet.append(["Артикул", "Кратность", "SKU", "Наименование", "Источник", "Прайс РТП", "Кратность Unitka", "Изменено"])
    labels = {"manual": "Вручную", "import": "Импорт", "rtp_price": "Прайс РТП", "unitka": "Unitka", "unknown": "Не задано"}
    for item in items:
        sheet.append([item["article"], item["pack_multiple"], ", ".join(item.get("skus") or []), item.get("product_name"), labels[item["source"]], item.get("rtp_price_pack_multiple"), item["unitka_pack_multiple"], item["updated_at"]])
    stream = BytesIO(); workbook.save(stream); workbook.close(); return stream.getvalue()
