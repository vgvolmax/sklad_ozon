"""Compact backend-built Excel exports with numeric finance cells."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _sheet(book, title, headers, rows, widths, *, percentages=(), freeze='A2'):
    sheet = book.create_sheet()
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='087F73')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[1].height = 32
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 42
        for index, cell in enumerate(row, 1):
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Article/name cannot become an Excel formula.
            if index in percentages:
                cell.number_format = '0.0%'
            elif isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='E9EEF2')
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:1'
    return sheet


def _workbook(title, headers, rows, widths, *, percentages=(), freeze='A2'):
    book = Workbook()
    book.remove(book.active)
    _sheet(book, title, headers, rows, widths, percentages=percentages, freeze=freeze)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def _number(value):
    return None if value is None else float(value)


def export_period_profit(report):
    book = Workbook(); book.remove(book.active)
    t = report['totals']
    _sheet(book, 'Итог', ['Показатель', 'Значение'], [
        ['Охват', 'Весь магазин · модель по текущей юнитке'],
        ['Период с', report['period']['from'].isoformat()],
        ['Период по', report['period']['to'].isoformat()],
        ['Режим', 'По заказам' if report['mode'] == 'orders' else 'По выкупам'],
        ['Количество, шт.', t['qty']], ['Модельная выручка, ₽', _number(t['revenue_model'])],
        ['Прибыль до общих расходов, ₽', _number(t['profit_before_common'])],
        ['Реклама за период, ₽', _number(t['advertising_total'])],
        ['Прочие общие расходы, ₽', _number(t['other_common_total'])],
        ['Прибыль после известных расходов, ₽', _number(t['profit_after_common'])],
        ['Маржа', _number(t['margin'])], ['Полнота', 'Частичный' if t['partial'] else 'Полный по известным данным'],
        ['Источник расходов', 'Финансовые начисления Ozon' if report['expenses_complete'] else 'Не загружены'], ['Загружено', report['loaded_at']],
        ['База юнитки', report.get('pricing_basis_id')], ['Снимок начислений', report.get('finance_snapshot_id')],
        ['SKU с полной базой', t['covered_sku_count']], ['SKU всего', t['sku_count']],
        ['Количество с полной базой', t['covered_qty']], ['SKU без полной базы', ', '.join(t['uncovered_skus'])],
        ['Прибыль выбранных SKU до общих расходов, ₽', _number(report['selected_profit_before_common'])]], [54, 66])
    _sheet(book, 'Товары', ['SKU', 'Артикул', 'Товар', 'Прибыль до рекламы, ₽ / шт.',
        'Количество, шт.', 'Вклад в прибыль, ₽', 'Цена продавца, ₽', 'Модельная выручка, ₽', 'Статус', 'Выкуплено, шт.', 'Возвращено, шт.'],
        [[p['sku'], p['article'], p['name'], _number(p['profit_per_unit_before_ads']), p['qty'],
          _number(p['profit']), _number(p['price']), _number(p['revenue']),
          'Неполный' if p['partial'] else 'Известен', p['purchased_qty'], p['returned_qty']] for p in report['products']], [24,22,50,28,22,26,26,28,22,22,22])
    _sheet(book, 'Расходы', ['Категория', 'Сумма всего магазина, ₽', 'Роль'],
        [[e['label'], _number(e['amount']), {'advertising':'Реклама · вычтена','additional_period_expense':'Общий расход · вычтен','already_in_unit_model':'Уже в юнитке','unclassified':'Не классифицирован · не вычтен'}[e['role']]] for e in report['expenses']], [60,30,38])
    stream = BytesIO(); book.save(stream); return stream.getvalue()


def export_buyouts(report):
    book = Workbook()
    book.remove(book.active)
    rows = [[p['sku'], p['article'], p['name'], p['purchased_qty'], p['returned_qty'],
             p['qty'] if p['quantity_known'] else None, _number(p['cost']),
             _number(p['cost_total']), _number(p['revenue']), _number(p['net_proceeds']),
             _number(p['known_expenses']), _number(p['profit'])] for p in report['products']]
    _sheet(book, 'Выкупы', ['SKU', 'Артикул', 'Товар', 'Выкуплено, шт.', 'Возвращено, шт.',
        'Выкупы минус возвраты, шт.', 'Текущая себестоимость, ₽ / шт.', 'Себестоимость всего, ₽',
        'Выручка продавца, ₽', 'Начислено за товары после расходов Ozon, ₽',
        'Расходы и корректировки в товарах, ₽', 'Прибыль по товарам, ₽'], rows,
        [24, 22, 54, 20, 20, 24, 26, 24, 24, 32, 32, 26])
    totals = report['totals']
    _sheet(book, 'Итог периода', ['Показатель', 'Значение'], [
        ['Период с', report['period']['from'].isoformat()], ['Период по', report['period']['to'].isoformat()],
        ['Себестоимость', 'Из текущей загрузки пользователя'],
        ['Прибыль по отфильтрованным товарам', _number(report['selected_totals']['profit'])],
        ['Прибыль по товарам всего магазина', _number(totals['profit_before_common'])],
        ['Общие расходы всего магазина', _number(totals['common_expenses'])],
        ['Реклама всего (уже учтена в расходах)', _number(totals['advertising_spend'])],
        ['Прибыль магазина после известных расходов', _number(totals['profit_after_known_expenses'])],
        ['Полнота расчёта', 'Частичный' if totals['partial'] else 'Полный по известным расходам'],
        ['SKU с рассчитанной прибылью', totals['covered_sku_count']],
        ['SKU всего', totals['sku_count']],
        ['Налоги вне Ozon', 'Не учтены'] ], [52, 44])
    _sheet(book, 'Расходы', ['Категория', 'Расходы всего магазина, ₽',
        'Уже в прибыли товаров, ₽', 'Вычитаются из общего итога, ₽'],
        [[e['label'], _number(e['amount']), _number(e['product_amount']), _number(e['common_amount'])]
         for e in report['expenses']], [60, 28, 30, 34])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def export_cost_prices(items):
    sources = {'manual': 'Ручная', 'import': 'Из юнитки'}
    rows = [[item['article'], item['product_name'], _number(item['cost']),
             sources[item['source']], item['updated_at']] for item in items]
    return _workbook('Себестоимость', ['Артикул', 'Товар', 'Себестоимость, ₽',
                                      'Источник', 'Обновлено (UTC)'], rows, [18, 60, 22, 18, 32])


def export_economics(report):
    drr_sources = {'real':'Реальный отчёт', 'plan':'План · отчёт отсутствует',
                   'zero_assumption':'Допущение 0 % · отчёт отсутствует', 'incomplete_report':'Отчёт неполный'}
    basis_sources = {'selected_period':'Выбранный период', 'loaded_history':'Загруженная история',
                     'scenario_worst_route':'Выбранные маршруты · худший вариант'}
    skus = set()
    rows = []
    for product in sorted(report['products'], key=lambda p: (p['article'], p['sku'])):
        article = product['article']
        sku = product['sku']
        if not sku or sku in skus:
            raise ValueError('Нужен уникальный SKU для каждой строки отчёта.')
        skus.add(sku)
        prices = product.get('buyer_prices') or {}
        rows.append([article, product['name'], _number(product['price']),
                     _number(product['planned_drr_rate']), _number(product['real_drr_rate']), _number(product['margin']),
                     _number(product['roi']), _number(report['target_margin']),
                     _number(product['target_price_all_routes']),
                     _number(product.get('applied_drr_rate')),
                     _number(product.get('commission_per_unit')),
                     (report.get('period') or {}).get('from'),
                     (report.get('period') or {}).get('to'), product['qty'], sku,
                     _number(prices.get('buyer_price_mean')), _number(prices.get('spp_mean')),
                     _number(prices.get('target_buyer_price')), prices.get('ordered_qty'),
                     prices.get('buyer_priced_qty'), prices.get('spp_priced_qty'),
                     'Полная' if prices.get('complete') else 'Неполная / отсутствует',
                     'Уточняются' if prices.get('pending') else 'По загруженным данным',
                     _number(product.get('cost')), product.get('cost_source'),
                     _number(product.get('profit_per_unit')), _number(product.get('profit_per_unit_before_ads')),
                     drr_sources.get(product.get('applied_drr_source'), 'Неизвестен'),
                     'Без продаж' if product.get('calculation_kind') == 'no_sales' else 'По истории',
                     basis_sources.get(product.get('basis_source'), 'История маршрутов'),
                     '; '.join(' → '.join(route) for route in product.get('calculation_routes', [])),
                     ' → '.join(product.get('limiting_route') or ()), product.get('price_action'),
                     'Полный' if product.get('pricing_complete') else 'Неполный'])
    return _workbook('Экономика', ['Артикул', 'Товар', 'Текущая цена, ₽',
        'ДРР по плану, %', 'Реальный ДРР, %', 'Маржа, %', 'ROI, %', 'Плановая маржа, %',
        'Необходимая цена, ₽', 'ДРР в расчёте, %', 'Комиссия Ozon, ₽ / шт.',
        'Период с', 'Период по', 'Доставлено, шт.', 'SKU', 'Средняя цена клиента, ₽',
        'Средний СПП, %', 'Цена клиента при цели, ₽', 'Заказано для средних, шт.',
        'Цена клиента известна, шт.', 'СПП известен, шт.', 'История заказов для цен', 'Статус цен',
        'Себестоимость, ₽ / шт.', 'Источник себестоимости', 'Прибыль, ₽ / шт.', 'Прибыль до рекламы, ₽ / шт.',
        'Источник ДРР в расчёте', 'Режим юнитки', 'Источник маршрутов', 'Расчётные маршруты',
        'Маршрут, ограничивающий цену', 'Действие с ценой', 'Расчёт юнитки'],
        rows, [18, 58, 22, 20, 20, 16, 16, 22, 24, 20, 24, 18, 18, 20, 24,
               26, 20, 28, 26, 28, 24, 28, 28, 26, 26, 26, 30, 40, 22, 38, 50, 50, 22, 22],
        percentages=(4, 5, 6, 7, 8, 10, 17), freeze='C2')
