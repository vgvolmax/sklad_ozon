"""Ozon Availability XLSX recommendation, scoped to exact current identities."""

from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
import re
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from openpyxl.utils.exceptions import InvalidFileException
try:
    from lxml.etree import XMLSyntaxError
except ImportError:  # openpyxl also supports the standard-library XML parser
    XMLSyntaxError = ParseError

from backend.domain.contracts import ImportDiagnostic, ImportResult, ReportMeta
from backend.ingestion._common import read_xlsx_tables, parse_decimal
from backend.ingestion.normalization import (normalize_cluster_label,
                                             normalize_seller_article_identity,
                                             normalize_text)
from backend.ozon.adapters.local_sale import RecommendedSupply, supply_period_for_days


_HEADER = re.compile(r"^рекомендуемая поставка,?\s*шт\s*на\s*(\d+)\s*дн", re.I)
_PERIOD = re.compile(r"период:\s*(\d{2}\.\d{2}\.\d{4})\s*[-–]\s*(\d{2}\.\d{2}\.\d{4})", re.I)
_UPDATED = re.compile(r"дата обновления:\s*(\d{2}\.\d{2}\.\d{4})\s+(\d{2}:\d{2})\s+мск", re.I)
_MAX_EXACT_EXCEL_INTEGER = (1 << 53) - 1


@dataclass(frozen=True, slots=True)
class RecommendationFileCheck:
    valid: bool
    usable_for_comparison: bool
    code: str | None = None
    message: str | None = None


def check_recommendation_import(imported: ImportResult, *, horizon_days: int,
                                source_as_of: date) -> RecommendationFileCheck:
    """One verdict shared by upload preflight and the final analysis boundary."""
    first_error = next((d for d in imported.diagnostics if d.severity == "error"), None)
    if first_error:
        return RecommendationFileCheck(False, False, "INVALID_RECOMMENDATION_FILE", first_error.message)
    if not imported.records:
        return RecommendationFileCheck(False, False, "EMPTY_RECOMMENDATION_FILE",
            "В отчёте нет подходящих строк для текущих SKU и кластеров Ozon.")
    if supply_period_for_days(horizon_days) is None:
        message = "Горизонт рекомендации Ozon не поддерживается: выберите 7, 14, 28 или 56 дней."
    elif horizon_days != imported.meta.recommendation_horizon_days:
        message = (f"Горизонт файла {imported.meta.recommendation_horizon_days} дней "
                   f"не совпадает с горизонтом расчёта {horizon_days} дней.")
    elif (datetime.fromisoformat(imported.meta.report_generated_at)
          .astimezone(timezone(timedelta(hours=3))).date() != source_as_of):
        message = ("Дата отчёта рекомендаций отличается от даты снимка API. "
                   "Обновите данные Ozon и выгрузите свежий XLSX.")
    else:
        return RecommendationFileCheck(True, True)
    return RecommendationFileCheck(True, False, "RECOMMENDATION_NOT_COMPARABLE", message)


def _metadata(data: bytes):
    from io import BytesIO
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        if sheet.calculate_dimension() == "A1:A1":
            sheet.reset_dimensions()
        labels = [normalize_text(row[0]) for row in sheet.iter_rows(max_row=5, max_col=1, values_only=True)]
    finally:
        workbook.close()
    period = next((_PERIOD.search(label) for label in labels if _PERIOD.search(label)), None)
    updated = next((_UPDATED.search(label) for label in labels if _UPDATED.search(label)), None)
    if not period or not updated:
        raise ValueError("Missing report period or update date")
    start, end = (datetime.strptime(value, "%d.%m.%Y").date() for value in period.groups())
    generated = datetime.strptime(" ".join(updated.groups()), "%d.%m.%Y %H:%M").replace(
        tzinfo=timezone(timedelta(hours=3))).astimezone(timezone.utc)
    if start > end:
        raise ValueError("Invalid report period")
    return start, end, generated


def import_recommendation_xlsx(data: bytes, filename: str, current_skus: frozenset[str] | set[str],
                               clusters: tuple) -> ImportResult[RecommendedSupply]:
    """Read only recommended quantities; never reuse the workbook's stock or sales."""
    from datetime import datetime, timezone

    imported = datetime.now(timezone.utc).isoformat()
    empty_meta = ReportMeta(filename, imported)
    if not data.startswith(b"PK"):
        return ImportResult((), (ImportDiagnostic("error", "INVALID_RECOMMENDATION_XLSX",
            "Выберите XLSX отчёт Ozon «Доступность товаров»."),), empty_meta)
    try:
        start, end, generated = _metadata(data)
        source = read_xlsx_tables(data, lambda names: (
            {"sku", "кластер", "схема продаж"} <= set(names)
            and any(_HEADER.match(name) for name in names)), read_only=True)
    except (ValueError, IndexError, KeyError, OSError, BadZipFile,
            InvalidFileException, ParseError, XMLSyntaxError):
        return ImportResult((), (ImportDiagnostic("error", "INVALID_RECOMMENDATION_XLSX",
            "Не удалось прочитать период и дату отчёта «Доступность товаров»."),), empty_meta)

    header = source.rows[0][1] if source.rows else {}
    horizons = {int(match.group(1)) for name in header
                if (match := _HEADER.match(name))}
    meta = ReportMeta(filename, imported, generated.isoformat(), start.isoformat(),
                      end.isoformat(), next(iter(horizons)) if len(horizons) == 1 else None)
    diagnostics = list(source.diagnostics)
    if len(horizons) != 1:
        diagnostics.append(ImportDiagnostic("error", "INVALID_RECOMMENDATION_HEADER",
            "Ожидается одна колонка рекомендации с указанным сроком поставки."))
        return ImportResult((), tuple(diagnostics), meta)
    quantity_column = next(name for name in header if _HEADER.match(name))
    by_name = {}
    for cluster in clusters:
        key = normalize_cluster_label(cluster.name).casefold()
        if key in by_name:
            diagnostics.append(ImportDiagnostic("error", "AMBIGUOUS_RECOMMENDATION_CLUSTER",
                "В каталоге Ozon повторяется название кластера."))
            return ImportResult((), tuple(diagnostics), meta)
        by_name[key] = normalize_cluster_label(cluster.name)

    accepted = {}
    source_row_by_identity = {}
    duplicates = set()
    for row_number, row in source.rows:
        sku = normalize_seller_article_identity(row.get("sku"))
        cluster = normalize_cluster_label(row.get("кластер"))
        if not sku or sku not in current_skus:
            diagnostics.append(ImportDiagnostic("warning", "UNKNOWN_RECOMMENDATION_SKU",
                "SKU отсутствует в текущем каталоге Ozon; строка исключена.", row_number, "sku"))
            continue
        canonical = by_name.get(cluster.casefold())
        if canonical is None:
            diagnostics.append(ImportDiagnostic("warning", "UNKNOWN_RECOMMENDATION_CLUSTER",
                f"Кластер «{cluster[:80]}» отсутствует в текущем каталоге Ozon; строка исключена.",
                row_number, "кластер"))
            continue
        if "FBO" not in {value.strip().upper() for value in normalize_text(row.get("схема продаж")).split(",")}:
            diagnostics.append(ImportDiagnostic("warning", "NON_FBO_RECOMMENDATION",
                "Рекомендация без схемы FBO исключена.", row_number, "схема продаж"))
            continue
        try:
            quantity = parse_decimal(row.get(quantity_column))
            if quantity < 0 or quantity != quantity.to_integral_value() or quantity > _MAX_EXACT_EXCEL_INTEGER:
                raise ValueError
        except ValueError:
            diagnostics.append(ImportDiagnostic("warning", "INVALID_RECOMMENDATION_QUANTITY",
                "Количество должно быть целым неотрицательным числом; строка исключена.",
                row_number, quantity_column))
            continue
        identity = sku, canonical
        if identity in accepted or identity in duplicates:
            if identity in accepted:
                diagnostics.append(ImportDiagnostic("warning", "DUPLICATE_RECOMMENDATION",
                    "Повтор SKU и кластера: все записи этой пары исключены.",
                    source_row_by_identity.pop(identity), "sku"))
            duplicates.add(identity)
            accepted.pop(identity, None)
            diagnostics.append(ImportDiagnostic("warning", "DUPLICATE_RECOMMENDATION",
                "Повтор SKU и кластера: все записи этой пары исключены.", row_number, "sku"))
            continue
        accepted[identity] = RecommendedSupply(sku, canonical, int(quantity))
        source_row_by_identity[identity] = row_number
    if not source.rows:
        diagnostics.append(ImportDiagnostic("error", "EMPTY_RECOMMENDATION_REPORT",
            "Не найдены строки рекомендаций Ozon."))
    return ImportResult(tuple(accepted.values()), tuple(diagnostics), meta)
