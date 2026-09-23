"""Single-open ingestion for Unitka economics and tariffs only."""

from dataclasses import dataclass
from io import BytesIO
from time import perf_counter
from typing import Callable

from openpyxl import load_workbook

from backend.domain.contracts import ImportResult, ProductEconomicsInput, ReportMeta, TariffRow
from ._common import _diag
from .normalization import normalize_header
from .product_economics import import_product_economics
from .tariffs import import_tariffs

_META_SHEET = "__sklad_ozon_meta"
_NATIVE_FORMAT = "sklad_ozon_unitka"
_NATIVE_VERSION = 1
_NATIVE_ECONOMICS_HEADERS = frozenset({"sku", "артикул", "себестоимость", "доступный остаток", "цена", "комиссия", "объём, л"})
_NATIVE_TARIFF_HEADERS = frozenset({"кластер отгрузки", "кластер доставки", "объём от", "объём до", "цена от", "цена до", "логистика"})


@dataclass(frozen=True, slots=True)
class UnitkaImportBundle:
    product_economics: ImportResult[ProductEconomicsInput]
    tariffs: ImportResult[TariffRow]
    source_format: str
    schema_version: int | None


def _native_metadata(workbook) -> tuple[str, int | None]:
    if _META_SHEET not in workbook.sheetnames:
        return "external", None
    values = {normalize_header(row[0]): row[1] for row in workbook[_META_SHEET].iter_rows(values_only=True)
              if len(row) >= 2 and row[0] is not None}
    if normalize_header(values.get("format")) != _NATIVE_FORMAT:
        return "native_invalid", None
    version = values.get("schema_version")
    return "native", version if isinstance(version, int) and not isinstance(version, bool) else None


def _has_exact_header(workbook, expected: frozenset[str]) -> bool:
    for sheet in workbook.worksheets:
        if sheet.title == _META_SHEET:
            continue
        for row in sheet.iter_rows(max_row=30, values_only=True):
            headers = frozenset(normalize_header(value) for value in row if normalize_header(value))
            if headers == expected:
                return True
    return False


def _invalid_native(report_context, code: str, message: str,
                    schema_version: int | None = None) -> UnitkaImportBundle:
    products = ImportResult((), (_diag(code, message),), report_context)
    return UnitkaImportBundle(products, ImportResult((), (), report_context), "native", schema_version)


def import_unitka_bundle(data: bytes, report_context: ReportMeta, *,
                         timing: Callable[[str, float, int | None], None] | None = None) -> UnitkaImportBundle:
    """Open Unitka once and preserve the standalone importers' parser contracts."""
    started = perf_counter()
    workbook = load_workbook(BytesIO(data), read_only=False, data_only=True)
    if timing:
        timing("unitka_open", perf_counter() - started, None)
    try:
        source_format, schema_version = _native_metadata(workbook)
        if source_format == "native_invalid":
            return _invalid_native(report_context, "INVALID_NATIVE_UNITKA_SCHEMA", "Native Unitka metadata is invalid.")
        if source_format == "native" and schema_version != _NATIVE_VERSION:
            return _invalid_native(report_context, "UNSUPPORTED_NATIVE_UNITKA_VERSION", "Native Unitka schema version is not supported.", schema_version)
        if source_format == "native" and not (_has_exact_header(workbook, _NATIVE_ECONOMICS_HEADERS) and _has_exact_header(workbook, _NATIVE_TARIFF_HEADERS)):
            return _invalid_native(report_context, "INVALID_NATIVE_UNITKA_SCHEMA", "Native Unitka columns do not match schema v1.", schema_version)
        started = perf_counter()
        tariffs = import_tariffs(data, report_context, workbook=workbook)
        if timing:
            timing("unitka_tariffs", perf_counter() - started, len(tariffs.records))
        started = perf_counter()
        products = import_product_economics(data, report_context, workbook=workbook)
        if timing:
            timing("unitka_economics", perf_counter() - started, len(products.records))
        return UnitkaImportBundle(products, tariffs, source_format, schema_version)
    finally:
        workbook.close()
