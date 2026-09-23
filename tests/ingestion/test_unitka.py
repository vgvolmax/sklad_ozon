from datetime import datetime, timezone
from io import BytesIO

import backend.ingestion.unitka as unitka_module
from openpyxl import Workbook
from backend.domain.contracts import ReportMeta
from backend.ingestion.product_economics import import_product_economics
from backend.ingestion.tariffs import import_tariffs
from tests.helpers.xlsx_fixtures import make_real_unitka

META = ReportMeta("unitka.xlsx", datetime.now(timezone.utc).isoformat())


def test_combined_unitka_matches_standalone_importers():
    data = make_real_unitka(product_rows=[["ART-1", "Товар", 100, 1000, "10%", 1]],
                            tariff_rows=[(0, "0-0,2 л", "Москва", "Казань", 18, 69)],
                            pack_rows=[["ART-1", "72/6"]])
    expected_products = import_product_economics(data, META)
    expected_tariffs = import_tariffs(data, META)
    bundle = unitka_module.import_unitka_bundle(data, META)
    assert bundle.product_economics == expected_products
    assert bundle.tariffs == expected_tariffs
    assert bundle.source_format == 'external'
    assert not hasattr(bundle, 'pack_multiplicity')


def test_combined_unitka_opens_workbook_once(monkeypatch):
    data = make_real_unitka()
    opens = 0
    original = unitka_module.load_workbook
    def counted(*args, **kwargs):
        nonlocal opens
        opens += 1
        return original(*args, **kwargs)
    monkeypatch.setattr(unitka_module, "load_workbook", counted)
    unitka_module.import_unitka_bundle(data, META)
    assert opens == 1


def test_external_unitka_ignores_invalid_supplier_packaging():
    data = make_real_unitka(pack_rows=[[40750.0, "garbage"]])
    bundle = unitka_module.import_unitka_bundle(data, META)
    diagnostics = bundle.product_economics.diagnostics + bundle.tariffs.diagnostics
    assert bundle.source_format == "external"
    assert "INVALID_PACK_MULTIPLICITY" not in {item.code for item in diagnostics}


def test_product_article_string_preserves_leading_zeroes():
    data = make_real_unitka(product_rows=[["00123", "Товар", 100, 1000, "10%", 1]])
    product = import_product_economics(data, META).records[0]
    assert product.article == "00123"


def native_unitka(version=1, *, omit_price=False):
    book = Workbook()
    products = book.active; products.title = "Товары"
    product_headers = ["SKU", "Артикул", "Себестоимость", "Доступный остаток",
                       "Цена", "Комиссия", "Объём, л"]
    if omit_price: product_headers.remove("Цена")
    products.append(product_headers)
    products.append(["SKU-1", "ART-1", 100, 5, "10%", 1] if omit_price else
                    ["SKU-1", "ART-1", 100, 5, 1000, "10%", 1])
    tariffs = book.create_sheet("Тарифы")
    tariffs.append(["Кластер отгрузки", "Кластер доставки", "Объём от", "Объём до",
                    "Цена от", "Цена до", "Логистика"])
    tariffs.append(["Москва", "Москва", 0, None, None, None, 69])
    meta = book.create_sheet("__sklad_ozon_meta")
    meta.append(["format", "sklad_ozon_unitka"]); meta.append(["schema_version", version])
    stream = BytesIO(); book.save(stream); book.close(); return stream.getvalue()


def test_native_unitka_requires_explicit_marker_and_accepts_strict_v1():
    bundle = unitka_module.import_unitka_bundle(native_unitka(), META)
    assert (bundle.source_format, bundle.schema_version) == ("native", 1)
    assert len(bundle.product_economics.records) == len(bundle.tariffs.records) == 1


def test_declared_native_unitka_fails_closed_for_schema_and_version():
    malformed = unitka_module.import_unitka_bundle(native_unitka(omit_price=True), META)
    unsupported = unitka_module.import_unitka_bundle(native_unitka(version=2), META)
    assert {item.code for item in malformed.product_economics.diagnostics} == {"INVALID_NATIVE_UNITKA_SCHEMA"}
    assert {item.code for item in unsupported.product_economics.diagnostics} == {"UNSUPPORTED_NATIVE_UNITKA_VERSION"}
    assert unsupported.schema_version == 2
