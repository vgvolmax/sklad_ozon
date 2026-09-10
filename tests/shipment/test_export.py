from io import BytesIO
from openpyxl import load_workbook
from backend.shipment.export import OZON_TEMPLATE_HEADERS,XLSX_MEDIA,render_export
from backend.ozon.draft_contracts import ValidationState
from backend.shipment.ranking import rank_outcomes
from .test_ranking import outcome,scenario

def test_exact_single_cluster_workbook():
    ranked,_=rank_outcomes((outcome("cs_a",ValidationState.ACCEPTED),),scenario())
    artifact=render_export(ranked[0],"sp_123456789")
    assert artifact.media_type==XLSX_MEDIA
    ws=load_workbook(BytesIO(artifact.content)).active
    assert tuple(ws.cell(1,x).value for x in range(1,4))==OZON_TEMPLATE_HEADERS
    assert ws["D1"].value is None and ws["B2"].value is None and ws["C2"].value==6
