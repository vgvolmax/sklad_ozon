"""Real loopback application and Seller client, with synthetic finance transport."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import socket
import tempfile
import threading
import time
from urllib.parse import urlparse

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect
import uvicorn

import backend.api as api
from backend.main import app
from backend.ozon.client import OzonClient, TransportResponse
from backend.ozon.contracts import OzonCredentials
from backend.ozon.vault import CredentialVault
from backend.security import LOCAL_SESSION_HEADER, current_local_session_token
from tests.api.test_analysis import _two_sku_analysis_files, _analysis_data, PRODUCT_HEADERS
from tests.helpers.xlsx_fixtures import make_xlsx
from tests.ozon.test_finance import posting, money

ROOT = Path(__file__).parents[2]
ARTIFACTS = ROOT / 'test-artifacts/economics-buyouts'


def main():
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as directory:
        api.PROJECT_PATH = Path(directory) / 'project.json'
        api.OZON_VAULT = CredentialVault(Path(directory) / 'vault.json')
        api.OZON_VAULT.setup(OzonCredentials('synthetic-account', 'synthetic-key'), 'test-password')
        api.ANALYSIS_STORE.clear(); api.FINANCE_STORE.clear()
        fail = [False]
        switch_account = [False]
        slow_read = [False]
        release_read = threading.Event()
        calls = []
        def transport(request, timeout):
            path = urlparse(request.full_url).path
            body = json.loads(request.data)
            calls.append((path, body))
            if path.endswith('/types'):
                payload = {'accrual_types': [{'id': 1, 'name': 'Реклама'}, {'id': 2, 'name': 'Кросс-докинг'}]}
            elif path.endswith('/by-day'):
                if slow_read[0]:
                    slow_read[0] = False
                    release_read.wait(timeout=8)
                if switch_account[0]:
                    switch_account[0] = False
                    api.OZON_VAULT.setup(OzonCredentials('other-account', 'other-key'), 'test-password')
                day = body['date']; rows = []
                if fail[0]:
                    payload = {'accruals': [posting(date=day,total_amount=money('NaN'))], 'last_id': ''}
                else:
                    if day == '2026-08-20':
                        first = posting(date=day); first['posting']['products'][0]['sku'] = 'SKU-1'
                        second = deepcopy(first); second['accrual_id'] = '2'; second['total_amount'] = money(600)
                        product = second['posting']['products'][0]; product['sku'] = 'SKU-2'
                        product['commission'].update(sale_amount=money(400), commission=money(300))
                        product['delivery']['total_accrued'] = money(-100)
                        rows = [first, second, {'accrual_id': 'ad', 'accrued_category': 'NON_ITEM', 'date': day,
                            'total_amount': money(-100), 'non_item_fee': {'type_id': 1, 'accrued': money(-100)}}]
                    if day == '2026-08-21':
                        returned = posting('return', date=day, total_amount=money(-600))
                        product = returned['posting']['products'][0]; product['sku'] = 'SKU-1'
                        product['commission'].update(sale_amount=money(-400), sale_price=money(-400), commission=money(-100))
                        product['delivery']['total_accrued'] = money(-100)
                        rows = [returned, {'accrual_id': 'crossdock', 'accrued_category': 'NON_ITEM', 'date': day,
                            'total_amount': money(-50), 'non_item_fee': {'type_id': 2, 'accrued': money(-50)}}]
                    payload = {'accruals': rows, 'last_id': ''}
            else:
                raise AssertionError(path)
            return TransportResponse(200, {}, json.dumps(payload).encode())
        api.OZON_CLIENT = OzonClient(api.OZON_VAULT, transport=transport, sleeper=lambda _: None)
        with TestClient(app, base_url='http://127.0.0.1', headers={LOCAL_SESSION_HEADER: current_local_session_token()}) as client:
            client.put('/api/project/cost-prices/ART-1', json={'cost': '999'})
            files = _two_sku_analysis_files()
            files['product_economics_file'] = ('products.xlsx', make_xlsx(headers=PRODUCT_HEADERS, rows=[
                ['SKU-1','ART-1',100,3,1000,'10%',1], ['SKU-2','ART-2',200,4,1200,'10%',1],
                ['NEW','NEW-ART',100,0,1000,'10%',1]]))
            response = client.post('/api/analysis', files=files, data=_analysis_data())
            assert response.status_code == 200, response.text
            snapshot = response.json()['snapshot']
        sock = socket.socket(); sock.bind(('127.0.0.1', 0))
        port = sock.getsockname()[1]
        server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
        thread = threading.Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True); thread.start()
        for _ in range(100):
            if server.started: break
            time.sleep(.05)
        assert server.started
        try:
            with sync_playwright() as p:
                browser = p.chromium.launch()
                context = browser.new_context(viewport={'width': 1440, 'height': 1000}, locale='ru-RU', accept_downloads=True,
                    permissions=['clipboard-read', 'clipboard-write'])
                page = context.new_page(); page.set_default_timeout(7000)
                errors, external = [], []
                page.on('pageerror', lambda error: errors.append(str(error)))
                page.on('request', lambda request: external.append(request.url) if not request.url.startswith(f'http://127.0.0.1:{port}/') else None)
                source = (ROOT/'frontend/assets/js/app.js').read_text().replace(
                    "if(root.document)document.addEventListener('DOMContentLoaded',S.boot);",
                    'S.__buyoutTest={setState(value){setState(value);}};')
                page.route('**/assets/js/app.js', lambda route: route.fulfill(body=source, content_type='text/javascript'))
                page.goto(f'http://127.0.0.1:{port}/')
                page.evaluate("snapshot=>{const S=SkladOzon;S.__buyoutTest.setState({...S.createInitialState(),section:'economics',snapshot});}", snapshot)
                expect(page.locator('[data-econ-row]')).to_have_count(3)
                expect(page.locator('#econ-target-form')).to_be_visible()
                page.locator('#econ-period-from').fill('2026-08-20')
                page.locator('#econ-period-to').fill('2026-08-21')
                page.locator('#econ-period-form button[type=submit]').click()
                expect(page.locator('[data-econ-period-summary]')).to_contain_text('2026-08-20')
                # A new SKU has explicit tariff routes and no history charts/SPP.
                new = page.locator('[data-econ-scenario-sku="NEW"]')
                expect(new).to_be_visible()
                expect(page.locator('[data-daily-sku="NEW"]')).to_have_count(0)
                new.locator('[data-econ-destination="Москва"]').check()
                expect(new.locator('[data-econ-target-price]')).not_to_contain_text('Не рассчитано')
                expect(new.locator('[data-econ-destination="Москва"]')).to_be_focused()
                new.screenshot(path=str(ARTIFACTS/'no-sales.png'))
                # Fallback policy changes the unit card, never the before-ad store basis.
                page.locator('[name=usePlan]').check()
                page.locator('#econ-target-form button[type=submit]').click()
                expect(page.locator('[data-econ-row="SKU-1"]')).to_contain_text('ДРР по плану')
                expect(page.locator('[data-econ-row="SKU-1"] td').nth(6)).to_contain_text('630')
                expect(page.locator('[data-econ-profit-before]')).to_have_text('1 426 ₽')
                page.locator('[name=usePlan]').uncheck()
                page.locator('#econ-target-form button[type=submit]').click()
                expect(page.locator('[data-econ-row="SKU-1"] td').nth(6)).to_contain_text('680')
                expect(page.locator('[data-econ-profit] [data-econ-calculation="buyouts"]')).to_be_visible()
                page.locator('[data-econ-calculation="buyouts"]').click()
                expect(page.locator('#econ-target-form')).to_be_visible()
                expect(page.locator('[data-econ-row]')).to_have_count(3)
                slow_read[0] = True
                page.locator('#buyout-load').click()
                expect(page.locator('#buyout-progress-detail')).to_contain_text('Ожидаем ответ Ozon')
                expect(page.locator('#buyout-progress-detail')).to_contain_text('20.08.2026')
                expect(page.locator('#buyout-progress-detail')).to_contain_text('Прошло: 1 с', timeout=3000)
                page.locator('[data-econ-calculation="orders"]').click()
                expect(page.locator('#econ-target-form')).to_be_visible()
                page.screenshot(path=str(ARTIFACTS/'slow-load.png'), full_page=True)
                page.locator('#buyout-cancel').click()
                expect(page.locator('#buyout-load')).to_be_enabled()
                expect(page.locator('#buyout-error')).to_contain_text('отменена')
                release_read.set()
                page.locator('#buyout-load').click()
                expect(page.locator('[data-buyout-final]')).to_have_text('1 276 ₽')
                expect(page.locator('[data-buyout-ads]')).to_have_text('100 ₽')
                # A failed mode refresh cannot turn retained orders into buyouts.
                page.route('**/api/economics/period/workspace', lambda route: route.fulfill(
                    status=500,content_type='application/json',body=json.dumps({'error':{'message':'Synthetic refresh failure'}})))
                page.locator('[data-econ-calculation="buyouts"]').click()
                expect(page.locator('#buyout-error')).to_contain_text('Synthetic refresh failure')
                expect(page.locator('[data-econ-profit-stale]')).to_contain_text('По заказам')
                expect(page.locator('[data-econ-profit] dl')).to_contain_text('Количество · заказы')
                expect(page.locator('[data-buyout-final]')).to_have_count(0)
                expect(page.locator('#buyout-export')).to_be_disabled()
                expect(page.locator('[data-econ-row]')).to_have_count(3)
                page.unroute('**/api/economics/period/workspace')
                page.locator('[data-econ-calculation="orders"]').click()
                expect(page.locator('[data-buyout-final]')).to_have_text('1 276 ₽')
                # Only this panel changes; the SKU nodes, draft and focus survive.
                page.locator('[data-econ-sku="SKU-1"]').click()
                expect(page.locator('.econ-clusters')).to_be_visible()
                draft = page.locator('[data-econ-drr="SKU-1"]')
                draft.fill('7')
                page.evaluate("()=>{window.__unitNode=document.querySelector('[data-econ-row]');document.querySelector('[data-econ-calculation=buyouts]').click();}")
                expect(page.locator('[data-buyout-final]')).to_have_text('1 956 ₽')
                expect(draft).to_have_value('7')
                expect(draft).to_be_focused()
                assert page.evaluate("window.__unitNode===document.querySelector('[data-econ-row]')")
                expect(page.locator('.econ-clusters')).to_be_visible()
                page.evaluate("document.querySelector('[data-econ-calculation=orders]').click()")
                expect(page.locator('[data-buyout-final]')).to_have_text('1 276 ₽')
                page.locator('[data-econ-close]').click()
                expect(page.locator('[data-econ-row]')).to_have_count(3)
                page.screenshot(path=str(ARTIFACTS/'desktop.png'), full_page=True)
                # Store expenses stay unchanged under the shared SKU filter.
                page.locator('#econ-search').fill('SKU-1')
                expect(page.locator('[data-econ-row]')).to_have_count(1)
                expect(page.locator('[data-buyout-final]')).to_have_text('1 276 ₽')
                with page.expect_download() as download:
                    page.locator('#buyout-export').click()
                book = load_workbook(BytesIO(Path(download.value.path()).read_bytes()))
                assert book.sheetnames == ['Итог','Товары','Расходы']
                assert book['Товары'].max_row == 2
                assert book['Товары']['A2'].value == 'SKU-1'
                expect(page.locator('#buyout-export')).to_be_enabled()
                page.locator('#econ-search-clear').click()
                fail[0] = True
                page.locator('#buyout-load').click()
                expect(page.locator('#buyout-error')).to_contain_text('Некорректная')
                expect(page.locator('[data-buyout-final]')).to_have_text('1 276 ₽')
                expect(page.locator('#econ-target-form')).to_be_visible()
                fail[0] = False
                page.locator('#buyout-load').click()
                expect(page.locator('#buyout-error')).to_be_hidden()
                page.locator('.buyout-expenses summary').click()
                expect(page.locator('.buyout-expenses')).to_contain_text('Кросс-докинг')
                for width, zoom, name in [(760,'1','narrow'),(760,'2','zoom-200')]:
                    page.set_viewport_size({'width':width,'height':800})
                    page.evaluate("z=>{document.documentElement.style.zoom=z;scrollTo(0,0)}", zoom)
                    page.screenshot(path=str(ARTIFACTS/(name+'.png')),full_page=True)
                    assert page.evaluate('document.documentElement.scrollWidth<=innerWidth')
                page.evaluate("document.documentElement.style.zoom='1'")
                # Re-entering the same FILES analysis revalidates the cabinet.
                api.OZON_VAULT.lock()
                page.evaluate("()=>{const S=SkladOzon;S.__buyoutTest.setState({...S.AppState});}")
                expect(page.locator('#buyout-error')).to_contain_text('Разблокируйте')
                expect(page.locator('[data-buyout-final]')).to_have_count(0)
                expect(page.locator('#econ-target-form')).to_be_visible()
                api.OZON_VAULT.unlock('test-password')
                page.locator('#buyout-load').click()
                expect(page.locator('[data-buyout-final]')).to_have_text('1 276 ₽')
                api.OZON_VAULT.lock()
                page.locator('#buyout-export').click()
                expect(page.locator('#buyout-error')).to_contain_text('Разблокируйте')
                expect(page.locator('[data-buyout-final]')).to_have_count(0)
                api.OZON_VAULT.unlock('test-password')
                page.locator('#buyout-load').click()
                expect(page.locator('[data-buyout-final]')).to_have_text('1 276 ₽')
                switch_account[0] = True
                page.locator('#buyout-load').click()
                expect(page.locator('#buyout-error')).to_be_visible()
                expect(page.locator('[data-buyout-final]')).to_have_count(0)
                expect(page.locator('#econ-target-form')).to_be_visible()
                assert not errors, errors
                assert not external, external
                browser.close()
        finally:
            server.should_exit = True; thread.join(timeout=10); sock.close()
        assert any(path.endswith('/by-day') for path, _ in calls)
    print('Economics buyout browser smoke passed')


if __name__ == '__main__':
    main()
