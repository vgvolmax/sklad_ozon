"""Real loopback API + production Economics UI, using synthetic files only."""
from io import BytesIO
import json
from pathlib import Path
import socket
import tempfile
import threading
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect
import uvicorn

import backend.api as api
from backend.main import app
from backend.security import LOCAL_SESSION_HEADER, current_local_session_token
from tests.api.test_analysis import _analysis_files, _analysis_data
from tests.economics.test_advertising import advertising_xlsx

ROOT = Path(__file__).parents[2]
ARTIFACTS = ROOT / 'test-artifacts/economics-ui'


def main():
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as directory:
        api.PROJECT_PATH = Path(directory) / 'project.json'
        api.ANALYSIS_STORE.clear()
        files = _analysis_files()
        files['orders_file'] = ('orders.csv', ('SKU;Количество;Цена продавца;Цена покупателя;Кластер отгрузки;Кластер доставки;Статус;Принят в обработку\n'
            'SKU-1;1;500;100;Москва;Москва;Доставлен;2026-09-01T10:00:00\n'
            'SKU-1;1;500;300;Москва;Москва;Доставлен;2026-09-01T12:00:00\n'
            'SKU-1;1;100;;Москва;Москва;Доставлен;2026-09-01T15:00:00\n'
            'SKU-1;1;2900;1450;Москва;Москва;Доставлен;2026-09-02T10:00:00\n').encode())
        with TestClient(app, base_url='http://127.0.0.1', headers={LOCAL_SESSION_HEADER:current_local_session_token()}) as client:
            response = client.post('/api/analysis', files=files, data=_analysis_data(
                as_of='2026-09-30', orders_period_from='2026-08-06', orders_period_to='2026-09-30'))
            assert response.status_code == 200, response.text
            snapshot = response.json()['snapshot']
        immutable_before = api.wire(api.ANALYSIS_STORE.latest())
        sock = socket.socket(); sock.bind(('127.0.0.1', 0))
        port = sock.getsockname()[1]
        server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
        thread = threading.Thread(target=server.run, kwargs={'sockets':[sock]}, daemon=True); thread.start()
        for _ in range(100):
            if server.started: break
            time.sleep(.05)
        assert server.started
        try:
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch()
                context = browser.new_context(viewport={'width':1440, 'height':960}, locale='ru-RU', accept_downloads=True)
                page = context.new_page()
                errors, external = [], []
                page.on('pageerror', lambda error: errors.append(str(error)))
                source = (ROOT / 'frontend/assets/js/app.js').read_text()
                hook = "if(root.document)document.addEventListener('DOMContentLoaded',S.boot);"
                assert hook in source
                source = source.replace(hook, "S.__browserEconomics={setState(value){setState(value);},getState(){return state;}};document.querySelectorAll('[data-nav]').forEach(link=>link.onclick=event=>{event.preventDefault();navigate(link.dataset.nav);});")
                page.route('**/assets/js/app.js', lambda route: route.fulfill(body=source, content_type='text/javascript'))
                def block_external(route):
                    if not route.request.url.startswith(f'http://127.0.0.1:{port}/'):
                        external.append(route.request.url); route.abort()
                    else: route.fallback()
                page.route('**/*', block_external)
                page.goto(f'http://127.0.0.1:{port}/')
                pending_series=[]
                page.route('**/api/economics/daily-series',lambda route:pending_series.append(route),times=1)
                page.evaluate("""snapshot=>{const S=SkladOzon,base=S.createInitialState();S.__browserEconomics.setState({...base,section:'economics',snapshot});}""", snapshot)
                expect(page.locator('.econ-real-drr')).to_have_text('Реальный n/a')
                expect(page.locator('.econ-drr-assumption')).to_have_text('В расчёте 0 %')
                assert 'Не рассчитано' not in page.locator('.econ-sku-row td').nth(6).inner_text()
                expect(page.locator('.econ-sku-row td').nth(4)).to_contain_text('₽ / шт.')
                actual=page.locator('.econ-sku-row td').nth(7).locator('.econ-fact')
                planned=page.locator('.econ-sku-row td').nth(7).locator('.econ-plan')
                expect(actual).to_contain_text('Факт')
                expect(planned).to_contain_text('План')
                assert actual.bounding_box()['y'] < planned.bounding_box()['y']
                assert actual.evaluate('(el)=>getComputedStyle(el).color') != planned.evaluate('(el)=>getComputedStyle(el).color')
                # Three charts share one calendar; first-day means are SPP 60%, buyer 200 RUB.
                panel=page.locator('[data-daily-sku="SKU-1"]')
                expect(panel).to_contain_text('Загружаем историю…')
                panel.locator('[data-daily-toggle]').click()
                assert pending_series
                pending_series.pop().fulfill(status=500,json={'error':{'message':'Проверка повтора истории'}})
                expect(panel).to_contain_text('Проверка повтора истории')
                expect(panel.locator('[data-daily-toggle]')).to_be_focused()
                panel.locator('[data-daily-retry]').click()
                expect(panel).to_contain_text('СПП 50 %–60 %')
                expect(panel.locator('[data-daily-toggle]')).to_be_focused()
                panel.locator('[data-daily-toggle]').click()
                assert panel.locator('.econ-daily-svg').count()==1
                expect(panel.locator('.econ-daily-svg')).to_have_css('height','110px')
                panel.locator('[data-daily-toggle]').click()
                expect(panel.locator('[data-daily-toggle]')).to_have_attribute('aria-expanded','true')
                plot=panel.locator('.econ-daily-plot')
                geometry=api.daily_series(api.ANALYSIS_STORE.latest().daily_order_evidence,'SKU-1')
                days=geometry['days'];index=next(i for i,d in enumerate(days) if d['day'].isoformat()=='2026-09-01')
                box=plot.bounding_box()
                page.mouse.move(box['x']+(64+920*(index+.5)/len(days))/1000*box['width'],box['y']+60)
                expect(panel.locator('.econ-daily-tooltip')).to_contain_text('СПП 60 %')
                expect(panel.locator('.econ-daily-tooltip')).to_contain_text('Заказы 3 шт.')
                expect(panel.locator('.econ-daily-tooltip')).to_contain_text('Цена покупателя 200 ₽')
                expect(panel.locator('.econ-daily-tooltip')).to_contain_text('СПП 2 / 3 шт.; покупатель 2 / 3 шт.')
                expect(panel.locator('.econ-daily-buyer-line')).to_have_count(1)
                plot.focus();page.keyboard.press('Home')
                expect(panel.locator('.econ-daily-tooltip')).to_contain_text('06.08.2026')
                expect(panel.locator('.econ-daily-tooltip')).to_contain_text('СПП n/a')
                expect(panel.locator('.econ-daily-tooltip')).to_contain_text('Заказы 0 шт.')
                page.keyboard.press('End')
                expect(panel.locator('.econ-daily-tooltip')).to_contain_text('30.09.2026')
                page.keyboard.press('Escape')
                expect(panel.locator('.econ-daily-tooltip')).to_be_hidden()
                page.screenshot(path=str(ARTIFACTS/'daily-expanded.png'),full_page=True)
                # Touch chooses the tapped day, rather than the previous selection.
                context_touch=browser.new_context(viewport={'width':1440,'height':960},has_touch=True,locale='ru-RU')
                touch=context_touch.new_page()
                touch.on('pageerror',lambda error:errors.append(str(error)))
                touch.route('**/assets/js/app.js',lambda route:route.fulfill(body=source,content_type='text/javascript'))
                touch.route('**/*',block_external)
                touch.goto(f'http://127.0.0.1:{port}/')
                touch.evaluate("""snapshot=>{const S=SkladOzon;S.__browserEconomics.setState({...S.createInitialState(),section:'economics',snapshot});}""",snapshot)
                touch_panel=touch.locator('[data-daily-sku="SKU-1"]')
                expect(touch_panel).to_contain_text('СПП 50 %–60 %')
                touch_panel.locator('[data-daily-toggle]').tap()
                touch_plot=touch_panel.locator('.econ-daily-plot');touch_plot.scroll_into_view_if_needed()
                touch_box=touch_plot.bounding_box()
                touch.touchscreen.tap(touch_box['x']+(64+920*(index+1.5)/len(days))/1000*touch_box['width'],touch_box['y']+60)
                expect(touch_panel.locator('.econ-daily-tooltip')).to_contain_text('02.09.2026')
                expect(touch_panel.locator('.econ-daily-tooltip')).to_contain_text('СПП 50 %')
                expect(touch_panel.locator('.econ-daily-tooltip')).to_contain_text('Цена покупателя 1\u00a0450 ₽')
                context_touch.close()
                panel.locator('[data-daily-toggle]').click()
                expect(panel.locator('[data-daily-toggle]')).to_have_attribute('aria-expanded','false')
                assert page.locator('[name=modelDrr]').count() == 0
                # Selection survives redraws and visiting the other screens.
                report = advertising_xlsx(rows=[['01.09.2026','SKU-1','Товар',100],['02.09.2026','SKU-1','Товар',300]])
                picker = page.locator('#econ-ads-files')
                picker.set_input_files([{'name':'campaign.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':report},
                                        {'name':'broken.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':b'broken'}])
                page.locator('[data-nav=data]').click()
                assert page.locator('#econ-ads-files').is_hidden()
                assert page.locator('#data-screen #econ-ads-files').count() == 0
                page.locator('[data-nav=economics]').click()
                expect(page.locator('.econ-advertising-list').first).to_contain_text('campaign.xlsx')
                page.route('**/api/economics/advertising/import',lambda route:route.fulfill(status=500,json={'error':{'message':'Проверка повтора после ошибки сохранения'}}),times=1)
                page.locator('#econ-ads-upload').click()
                expect(page.locator('.econ-advertising')).to_contain_text('Проверка повтора после ошибки сохранения')
                assert not api.PROJECT_PATH.with_name('advertising.json').exists()
                page.locator('#econ-ads-upload').click()
                expect(page.locator('.econ-real-drr')).to_have_text('Реальный 10 %')
                expect(page.locator('.econ-ads-matches').first).to_contain_text('Сопоставлено по SKU: 1')
                page.locator('.econ-ads-matches summary').first.click()
                expect(page.locator('.econ-ads-matches').first).to_contain_text('SKU SKU-1 · ART-1')
                expect(page.locator('.econ-advertising')).to_contain_text('Не удалось прочитать XLSX')
                assert not page.evaluate('SkladOzon.__browserEconomics.getState().staleSnapshot')
                assert api.wire(api.ANALYSIS_STORE.latest()) == immutable_before
                page.screenshot(path=str(ARTIFACTS / 'desktop.png'), full_page=True)
                # Remove only the failed file; repeat is idempotent.
                page.locator('[data-econ-ads-remove-file="1"]').click()
                picker.set_input_files({'name':'repeat.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':report})
                before = api.PROJECT_PATH.with_name('advertising.json').read_bytes()
                page.locator('#econ-ads-upload').click()
                expect(page.locator('.econ-advertising')).to_contain_text('Уже загружено')
                expect(page.locator('#econ-ads-files')).to_be_enabled()
                assert api.PROJECT_PATH.with_name('advertising.json').read_bytes() == before
                # A corrected same-campaign export replaces; another campaign adds.
                corrected = advertising_xlsx(rows=[['01.09.2026','SKU-1','Товар',200],['02.09.2026','SKU-1','Товар',500]])
                extra = advertising_xlsx(campaign='222', rows=[['01.09.2026','SKU-1','Товар',100],['02.09.2026','SKU-1','Товар',100]])
                picker.set_input_files([{'name':'corrected.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':corrected},
                                        {'name':'extra.xlsx','mimeType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','buffer':extra}])
                page.locator('#econ-ads-upload').click()
                expect(page.locator('.econ-real-drr')).to_have_text('Реальный 22,5 %')
                # Actual numeric rate in downloaded report, unaffected by filters.
                page.locator('#econ-search').fill('нет такого товара')
                with page.expect_download() as download:
                    page.locator('#econ-export').click()
                book = load_workbook(BytesIO(Path(download.value.path()).read_bytes()))
                sheet = book.active
                assert sheet.max_row == 2 and sheet.cell(2,5).value == .225
                page.locator('#econ-search-clear').click()
                # Shared confirmation cancel/Escape is harmless; confirm recalculates.
                page.locator('.econ-advertising-campaigns summary').click()
                page.locator('[data-econ-ads-delete="222"]').click()
                page.keyboard.press('Escape')
                expect(page.locator('.econ-real-drr')).to_have_text('Реальный 22,5 %')
                page.locator('[data-econ-ads-delete="222"]').click()
                page.locator('[data-dialog-confirm]').click()
                expect(page.locator('.econ-real-drr')).to_have_text('Реальный 17,5 %')
                page.locator('.econ-advertising-campaigns summary').click()
                page.locator('[data-econ-ads-delete="123"]').click()
                page.locator('[data-dialog-confirm]').click()
                expect(page.locator('.econ-real-drr')).to_have_text('Реальный n/a')
                expect(page.locator('.econ-drr-assumption')).to_have_text('В расчёте 0 %')
                # Narrow window, horizontal table overflow stays local.
                page.locator('#econ-ads-clear-completed').click()
                assert page.locator('[data-econ-ads-remove-file]').count() == 0
                page.set_viewport_size({'width':720,'height':900})
                page.evaluate('window.scrollTo(0,0)')
                assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
                panel.locator('[data-daily-toggle]').click()
                plot=panel.locator('.econ-daily-plot');box=plot.bounding_box()
                page.mouse.move(min(690,box['x']+box['width']*.48),box['y']+40)
                popup=panel.locator('.econ-daily-tooltip');expect(popup).to_be_visible()
                popup_box=popup.bounding_box()
                assert popup_box['x']>=0 and popup_box['x']+popup_box['width']<=720
                page.screenshot(path=str(ARTIFACTS / 'narrow.png'), full_page=True)
                assert not errors, errors
                assert not external, external
                browser.close()
                print('Economics browser: DRR-zero, commission, daily price/SPP averages and coverage, actual/plan colors, delayed error/retry/focus, hover, keyboard, touch, batch, errors, navigation, duplicate, correction, multi-campaign, export, deletion, narrow tooltip passed; no JS errors or external requests.')
        finally:
            server.should_exit = True; thread.join(timeout=5); sock.close()


if __name__ == '__main__': main()
