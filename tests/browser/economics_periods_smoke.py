"""Selected Economics intervals and lower-article edits in the production UI."""
from io import BytesIO
from pathlib import Path
import socket
import tempfile
import threading
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect
import uvicorn

import backend.api as api
from backend.main import app
from backend.security import LOCAL_SESSION_HEADER, current_local_session_token
from tests.api.test_analysis import _analysis_files, _analysis_data, PRODUCT_HEADERS
from tests.helpers.xlsx_fixtures import make_xlsx

ROOT = Path(__file__).parents[2]
ARTIFACTS = ROOT / 'test-artifacts/economics-periods'


def main():
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as directory:
        api.PROJECT_PATH = Path(directory) / 'project.json'
        api.ANALYSIS_STORE.clear()
        files = _analysis_files()
        files['product_economics_file'] = ('products.xlsx', make_xlsx(headers=PRODUCT_HEADERS,
            rows=[[f'SKU-{i}', f'ART-{i}', 100, 20, 1000, '10%', 1] for i in range(1, 4)]))
        files['orders_file'] = ('orders.csv', ('SKU;Количество;Цена продавца;Цена покупателя;Кластер отгрузки;Кластер доставки;Статус;Принят в обработку\n' +
            ''.join(f'SKU-{i};1;1000;200;Москва;Москва;Доставлен;2026-09-01T10:00:00\n'
                    f'SKU-{i};3;1000;800;Москва;Москва;Доставлен;2026-09-02T12:00:00\n'
                    f'SKU-{i};2;1000;500;Москва;Москва;Доставлен;2026-09-07T15:00:00\n'
                    f'SKU-{i};7;1000;600;Москва;Москва;Доставлен;2026-09-30T10:00:00\n'
                    for i in range(1, 4))).encode())
        with TestClient(app, base_url='http://127.0.0.1', headers={LOCAL_SESSION_HEADER: current_local_session_token()}) as client:
            response = client.post('/api/analysis', files=files, data=_analysis_data(
                as_of='2026-09-30', orders_period_from='2026-09-01', orders_period_to='2026-09-30'))
            assert response.status_code == 200, response.text
            snapshot = response.json()['snapshot']
        before = api.wire(api.ANALYSIS_STORE.latest())
        sock = socket.socket(); sock.bind(('127.0.0.1', 0))
        port = sock.getsockname()[1]
        server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
        thread = threading.Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True); thread.start()
        for _ in range(100):
            if server.started: break
            time.sleep(.05)
        try:
            with sync_playwright() as p:
                browser = p.chromium.launch()
                context = browser.new_context(viewport={'width': 1440, 'height': 900}, locale='ru-RU', accept_downloads=True)
                page = context.new_page(); page.set_default_timeout(5000)
                errors = []; page.on('pageerror', lambda error: errors.append(str(error)))
                source = (ROOT / 'frontend/assets/js/app.js').read_text()
                hook = "if(root.document)document.addEventListener('DOMContentLoaded',S.boot);"
                source = source.replace(hook, "S.__periodTest={setState(value){setState(value);}};")
                page.route('**/assets/js/app.js', lambda route: route.fulfill(body=source, content_type='text/javascript'))
                page.goto(f'http://127.0.0.1:{port}/')
                page.evaluate("snapshot=>{const S=SkladOzon;S.__periodTest.setState({...S.createInitialState(),section:'economics',snapshot});}", snapshot)
                expect(page.locator('.econ-sku-row')).to_have_count(3)
                row = page.locator('[data-econ-row="SKU-3"]')
                expect(row.locator('[data-econ-buyer-price]')).to_have_text('600 ₽')
                expect(row.locator('[data-econ-price-cell]')).to_contain_text('При среднем СПП 40 %')
                fact_color = row.locator('[data-econ-buyer-price]').evaluate('el=>getComputedStyle(el).color')
                plan_color = row.locator('[data-econ-target-buyer-price]').evaluate('el=>getComputedStyle(el).color')
                assert fact_color != plan_color
                # Capture the final editable article after returning from earlier disclosures.
                for sku in ('SKU-1', 'SKU-2', 'SKU-3'):
                    page.locator(f'[data-econ-sku="{sku}"]').click()
                    page.locator(f'[data-daily-sku="{sku}"] [data-daily-toggle]').click()
                    if sku != 'SKU-3':
                        page.locator('[data-econ-close]').click()
                drr = page.locator('[data-econ-drr="SKU-3"]')
                drr.scroll_into_view_if_needed(); drr.focus()
                position = drr.bounding_box()['y']; scroll = page.evaluate('scrollY')
                table_scroll = page.locator('.econ-table-scroll').evaluate('el=>el.scrollTop')
                assert scroll > 500
                pending = []
                page.route('**/api/economics/workspace', lambda route: pending.append(route), times=1)
                drr.fill('11'); drr.dispatch_event('change')
                expect(drr).to_be_focused()
                assert abs(drr.bounding_box()['y'] - position) < 3, (position, drr.bounding_box()['y'], scroll, page.evaluate('scrollY'))
                assert abs(page.locator('.econ-table-scroll').evaluate('el=>el.scrollTop') - table_scroll) < 3
                # Move while awaiting the response: preserve the user's new position.
                page.mouse.wheel(0, 100); page.wait_for_timeout(250)
                shifted = page.evaluate('scrollY')
                shifted_table = page.locator('.econ-table-scroll').evaluate('el=>el.scrollTop')
                pending.pop().continue_()
                expect(page.locator('.econ-page')).to_have_attribute('aria-busy', 'false')
                assert abs(page.evaluate('scrollY') - shifted) < 3, (shifted,page.evaluate('scrollY'),drr.bounding_box(),page.locator('[data-econ-profit]').bounding_box())
                assert abs(page.locator('.econ-table-scroll').evaluate('el=>el.scrollTop') - shifted_table) < 3
                expect(page.locator('[data-econ-sku="SKU-3"]')).to_have_attribute('aria-expanded', 'true')
                assert page.locator('[data-econ-row="SKU-3"]').evaluate("el=>parseFloat(getComputedStyle(el.cells[0]).borderTopWidth)") >= 3
                page.locator('#econ-period-from').fill('2026-09-02')
                page.locator('#econ-period-to').fill('2026-09-07')
                page.locator('#econ-period-apply').click()
                expect(page.locator('.econ-page')).to_have_attribute('aria-busy', 'false')
                row = page.locator('[data-econ-row="SKU-3"]')
                expect(row.locator('td').nth(1)).to_contain_text('5 шт.')
                expect(row.locator('[data-econ-buyer-price]')).to_have_text('680 ₽')
                expect(row.locator('[data-econ-price-cell]')).to_contain_text('При среднем СПП 32 %')
                expect(page.locator('[data-econ-period-summary]')).to_contain_text('2026-09-02–2026-09-07')
                page.locator('[data-econ-granularity="week"]').click()
                panel = page.locator('[data-daily-sku="SKU-3"]')
                expect(panel.locator('.econ-daily-period')).to_contain_text('По неделям')
                plot = panel.locator('.econ-daily-plot'); plot.focus(); plot.press('Home')
                expect(panel.locator('.econ-daily-tooltip')).to_contain_text('02.09.2026–06.09.2026')
                expect(panel.locator('.econ-daily-tooltip')).to_contain_text('СПП 20 %')
                expect(panel.locator('.econ-daily-tooltip')).to_contain_text('Цена покупателя 800 ₽')
                expect(panel.locator('.econ-daily-tooltip')).to_contain_text('Заказы 3 шт.')
                with page.expect_download() as downloaded:
                    page.locator('#econ-export').click()
                sheet = load_workbook(BytesIO(Path(downloaded.value.path()).read_bytes())).active
                assert sheet['L2'].value.strftime('%Y-%m-%d') == '2026-09-02'
                assert sheet['M2'].value.strftime('%Y-%m-%d') == '2026-09-07'
                assert sheet['N2'].value == 5
                assert sheet['P2'].value == 680 and sheet['Q2'].value == .32
                assert round(sheet['I2'].value * .68, 2) == sheet['R2'].value
                sku3_price = next(cells[17].value for cells in sheet.iter_rows(min_row=2)
                                  if cells[14].value == 'SKU-3')
                expect(row.locator('[data-econ-target-buyer-price]')).to_have_text(
                    page.evaluate("value=>SkladOzon.EconomicsWorkspace.formatMoney(value)", sku3_price))
                page.screenshot(path=str(ARTIFACTS / 'weekly-selected.png'), full_page=True)
                page.locator('#econ-period-reset').click()
                expect(page.locator('[data-econ-row="SKU-3"]').locator('td').nth(1)).to_contain_text('13 шт.')
                expect(row.locator('[data-econ-buyer-price]')).to_have_text('600 ₽')
                # A pending background-price response is refreshed automatically.
                def pending_prices(route):
                    response = route.fetch(); payload = response.json()
                    for series in payload['series'].values():
                        series['price_pending'] = True
                        for day in series['days']:
                            day['spp'] = day['buyer_price_mean'] = None
                            day['spp_priced_qty'] = day['buyer_priced_qty'] = 0
                    route.fulfill(response=response, json=payload)
                page.route('**/api/economics/daily-series', pending_prices, times=1)
                page.locator('[data-econ-granularity="day"]').click()
                expect(panel.locator('.econ-daily-buyer-line')).to_have_count(0)
                expect(panel.locator('.econ-daily-buyer-line')).to_have_count(3, timeout=12000)
                # Price means arrive independently: patch only price cells, never
                # redraw active forms or apply a response from an older request.
                def pending_means(route):
                    response = route.fetch(); payload = response.json()
                    payload['workspace']['prices_pending'] = True
                    for product in payload['workspace']['products']:
                        prices = product['buyer_prices']
                        prices.update(buyer_price_mean=None, target_buyer_price=None,
                                      spp_mean=None, buyer_priced_qty=0, spp_priced_qty=0, pending=True)
                    route.fulfill(response=response, json=payload)
                page.route('**/api/economics/workspace', pending_means, times=1)
                page.locator('#econ-period-reset').click()
                expect(row.locator('[data-econ-buyer-price]')).to_have_text('Не рассчитано')
                mean_requests = []
                page.route('**/api/economics/workspace', lambda route: mean_requests.append(route), times=1)
                margin = page.locator('#econ-target-form [name="margin"]')
                margin.fill('23'); margin.focus()
                page.wait_for_function('document.querySelector("[data-econ-price-status]").textContent.includes("уточняются")')
                for _ in range(100):
                    if mean_requests: break
                    page.wait_for_timeout(50)
                assert mean_requests
                mean_requests.pop().continue_()
                expect(row.locator('[data-econ-buyer-price]')).to_have_text('600 ₽')
                expect(margin).to_have_value('23'); expect(margin).to_be_focused()
                expect(page.locator('[data-econ-price-status]')).to_be_hidden()
                # A failed background refresh retains known data and a local retry.
                page.route('**/api/economics/workspace', pending_means, times=1)
                page.locator('#econ-period-reset').click()
                page.route('**/api/economics/workspace', lambda route: route.fulfill(
                    status=503, json={'error': {'message': 'Повторите уточнение цен'}}), times=1)
                expect(page.locator('[data-econ-price-retry]')).to_be_visible(timeout=8000)
                expect(page.locator('.econ-sku-row')).to_have_count(1)
                page.locator('[data-econ-price-retry]').click()
                expect(row.locator('[data-econ-buyer-price]')).to_have_text('600 ₽')
                expect(page.locator('[data-econ-price-status]')).to_be_hidden()
                # A chart-only redraw may start a newer price request while an
                # older one is slow. The older mean must not replace the new one.
                page.route('**/api/economics/workspace', pending_means, times=1)
                page.locator('#econ-period-reset').click()
                overlap = []
                page.route('**/api/economics/workspace', lambda route: overlap.append(route), times=1)
                for _ in range(100):
                    if overlap: break
                    page.wait_for_timeout(50)
                assert overlap
                page.locator('[data-econ-granularity="week"]').click()
                expect(row.locator('[data-econ-buyer-price]')).to_have_text('600 ₽', timeout=8000)
                older = overlap.pop()
                response = older.fetch(); payload = response.json()
                payload['workspace']['prices_pending'] = True
                for product in payload['workspace']['products']:
                    product['buyer_prices'].update(buyer_price_mean='40', pending=True)
                older.fulfill(response=response, json=payload)
                page.wait_for_timeout(150)
                assert row.locator('[data-econ-buyer-price]').inner_text() == '600 ₽'
                expect(page.locator('[data-econ-price-status]')).to_be_hidden()
                # Superseded refresh cannot replace the newly selected interval.
                page.route('**/api/economics/workspace', pending_means, times=1)
                page.locator('#econ-period-reset').click()
                late = []
                page.route('**/api/economics/workspace', lambda route: late.append(route), times=1)
                for _ in range(100):
                    if late: break
                    page.wait_for_timeout(50)
                assert late
                page.locator('#econ-period-from').fill('2026-09-02')
                page.locator('#econ-period-to').fill('2026-09-07')
                page.locator('#econ-period-apply').click()
                expect(row.locator('[data-econ-buyer-price]')).to_have_text('680 ₽')
                late.pop().continue_(); page.wait_for_timeout(150)
                expect(row.locator('[data-econ-buyer-price]')).to_have_text('680 ₽')
                page.screenshot(path=str(ARTIFACTS / 'buyer-prices-desktop.png'), full_page=True)
                page.set_viewport_size({'width': 390, 'height': 844})
                page.locator('#econ-period-from').scroll_into_view_if_needed()
                assert page.evaluate('document.documentElement.scrollWidth <= innerWidth + 1')
                page.screenshot(path=str(ARTIFACTS / 'narrow-period.png'))
                page.locator('[data-econ-price-cell]').evaluate('el=>el.scrollIntoView({inline:"center",block:"center"})')
                expect(row.locator('[data-econ-buyer-price]')).to_be_visible()
                expect(row.locator('[data-econ-target-buyer-price]')).to_be_visible()
                header = page.locator('.app-header').bounding_box()
                close_bar = page.locator('.econ-focus-bar').bounding_box()
                assert abs(close_bar['y'] - (header['y'] + header['height'])) < 3, (header, close_bar)
                page.screenshot(path=str(ARTIFACTS / 'buyer-prices-narrow.png'))
                page.set_viewport_size({'width': 720, 'height': 900})
                page.evaluate('document.documentElement.style.zoom="2";dispatchEvent(new Event("resize"))')
                page.locator('[data-econ-price-cell]').evaluate('el=>el.scrollIntoView({inline:"center",block:"center"})')
                page.screenshot(path=str(ARTIFACTS / 'buyer-prices-zoom.png'))
                assert page.evaluate('document.documentElement.scrollWidth <= innerWidth + 1'), page.evaluate(
                    "({width:innerWidth,scroll:document.documentElement.scrollWidth,panels:[...document.querySelectorAll('.econ-page,.econ-page>*')].map(n=>[n.className,n.getBoundingClientRect().width,getComputedStyle(n).minWidth]),overflow:[...document.querySelectorAll('.econ-page *')].filter(n=>n.getBoundingClientRect().right>innerWidth&&!n.closest('.econ-table-scroll')).map(n=>[n.tagName,n.className,n.getBoundingClientRect().right,getComputedStyle(n).minWidth,getComputedStyle(n).whiteSpace]).slice(0,30)})")
                expect(row.locator('[data-econ-target-buyer-price]')).to_be_visible()
                header = page.locator('.app-header').bounding_box()
                close_bar = page.locator('.econ-focus-bar').bounding_box()
                assert abs(close_bar['y'] - (header['y'] + header['height'])) < 3, (header, close_bar)
                assert not errors, errors
                assert api.wire(api.ANALYSIS_STORE.latest()) == before
                browser.close()
                print('Economics periods browser: stable edit/scroll, selected quantities and client means, target client price/XLSX, weighted weeks, background price retry, overlapping/stale responses, input preservation, separators, narrow and 200% layout passed.')
        finally:
            server.should_exit = True; thread.join(timeout=10); sock.close()


if __name__ == '__main__':
    main()
