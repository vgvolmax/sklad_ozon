from dataclasses import replace
from io import BytesIO
import json

import pytest
from openpyxl import Workbook, load_workbook

from backend.pack_multiplicity import (
    build_effective_pack_evidence, export_xlsx, parse_import_xlsx, reset_override, resolve_pack_multiplicity,
    pack_multiplicity_fingerprint, set_override,
    validate_pack_multiple,
)
from backend.project import PackMultiplicityRecord, Project, load_project, save_project_atomic


def workbook(rows):
    book=Workbook(); sheet=book.active
    for row in rows: sheet.append(row)
    stream=BytesIO(); book.save(stream); book.close(); return stream.getvalue()


def test_effective_precedence_reset_and_origins():
    base=PackMultiplicityRecord(unitka_pack_multiple=20)
    assert (resolve_pack_multiplicity(base).pack_multiple, resolve_pack_multiplicity(base).source) == (None, 'unknown')
    project,article=set_override(Project(pack_multiplicity={'17261':base}),17261.0,40,'manual',updated_at='now')
    resolved=resolve_pack_multiplicity(project.pack_multiplicity[article])
    assert (resolved.pack_multiple,resolved.source)==(40,'manual')
    project,_=set_override(project,'17261',50,'import',updated_at='later')
    assert resolve_pack_multiplicity(project.pack_multiplicity['17261']).source == 'import'
    project,_=reset_override(project,'17261')
    assert resolve_pack_multiplicity(project.pack_multiplicity['17261']).pack_multiple is None
    assert resolve_pack_multiplicity(None).source == 'unknown'


def test_rtp_price_precedence_and_reset_to_rtp():
    base=PackMultiplicityRecord(unitka_pack_multiple=10,
        rtp_price_pack_multiple=15,rtp_price_updated_at='price')
    resolved=resolve_pack_multiplicity(base)
    assert (resolved.pack_multiple,resolved.source)==(15,'rtp_price')
    project,_=set_override(Project(pack_multiplicity={'A':base}),'A',20,'manual',updated_at='now')
    assert resolve_pack_multiplicity(project.pack_multiplicity['A']).pack_multiple==20
    project,_=reset_override(project,'A')
    assert (resolve_pack_multiplicity(project.pack_multiplicity['A']).pack_multiple,
            resolve_pack_multiplicity(project.pack_multiplicity['A']).source)==(15,'rtp_price')


def test_legacy_unitka_is_ignored_by_effective_evidence():
    project=Project(pack_multiplicity={
        '17261':PackMultiplicityRecord(unitka_pack_multiple=50),
        '28202':PackMultiplicityRecord(unitka_pack_multiple=50, rtp_price_pack_multiple=40, rtp_price_updated_at='now'),
        '39439':PackMultiplicityRecord(unitka_pack_multiple=50, rtp_price_pack_multiple=40, rtp_price_updated_at='now', override_pack_multiple=20, override_origin='manual', override_updated_at='now'),
    })
    effective={item.article:item for item in build_effective_pack_evidence(project)}
    assert (effective['17261'].pack_multiple,effective['17261'].source)==(None,'unknown')
    assert (effective['28202'].pack_multiple,effective['28202'].source)==(40,'rtp_price')
    assert (effective['39439'].pack_multiple,effective['39439'].source)==(20,'manual')


def test_pack_fingerprint_is_deterministic_and_pack_specific():
    left=Project(pack_multiplicity={
        'B':PackMultiplicityRecord(20), 'A':PackMultiplicityRecord(10),
    },manual_cluster_mappings={'old':'new'})
    reordered=Project(pack_multiplicity={
        'A':PackMultiplicityRecord(10), 'B':PackMultiplicityRecord(20),
    })
    changed=Project(pack_multiplicity={
        'A':PackMultiplicityRecord(rtp_price_pack_multiple=10, rtp_price_updated_at='now'), 'B':PackMultiplicityRecord(rtp_price_pack_multiple=50, rtp_price_updated_at='now'),
    })
    assert pack_multiplicity_fingerprint(left) == pack_multiplicity_fingerprint(reordered)
    assert pack_multiplicity_fingerprint(left) != pack_multiplicity_fingerprint(changed)

def test_pack_fingerprint_ignores_display_timestamps():
    left=Project(pack_multiplicity={'A':PackMultiplicityRecord(
        rtp_price_pack_multiple=15,rtp_price_updated_at='old',
        override_pack_multiple=20,override_origin='manual',override_updated_at='old')})
    right=Project(pack_multiplicity={'A':replace(left.pack_multiplicity['A'],
        rtp_price_updated_at='new',override_updated_at='new')})
    assert pack_multiplicity_fingerprint(left)==pack_multiplicity_fingerprint(right)

@pytest.mark.parametrize('value', [1,50])
def test_valid_pack_multiple(value): assert validate_pack_multiple(value)==value
@pytest.mark.parametrize('value', [0,-1,1.5,True,'50 шт.',float('nan')])
def test_invalid_pack_multiple(value):
    with pytest.raises(ValueError): validate_pack_multiple(value)

def test_v1_migrates_and_v2_round_trip_persists(tmp_path):
    path=tmp_path/'project.json'; save_project_atomic(path,Project())
    payload=json.loads(path.read_text()); payload['schema_version']=1; payload.pop('pack_multiplicity'); path.write_text(json.dumps(payload))
    assert load_project(path).pack_multiplicity=={}
    project=Project(pack_multiplicity={'17261':PackMultiplicityRecord(unitka_pack_multiple=20,override_pack_multiple=50,override_origin='manual',override_updated_at='now')})
    save_project_atomic(path,project); assert load_project(path)==project

def test_xlsx_mixed_duplicate_upsert_and_reimportable_export():
    parsed=parse_import_xlsx(workbook([['Артикул','Кратность'],[17261.0,50],['BAD',0],['39439',8],['39439',9]])); values,errors=parsed.values,parsed.diagnostics
    assert values=={'17261':50}
    assert {e.code for e in errors}=={'INVALID_PACK_MULTIPLICITY','DUPLICATE_ARTICLE'}
    data=export_xlsx([{'article':'17261','pack_multiple':50,'source':'manual','rtp_price_pack_multiple':20,'override_pack_multiple':50,'updated_at':'now','skus':['1'],'product_name':'Товар'}])
    parsed=parse_import_xlsx(data); imported,errors=parsed.values,parsed.diagnostics
    assert imported=={'17261':50} and not errors
    sheet=load_workbook(BytesIO(data),read_only=True).active
    assert next(sheet.values)[:2]==('Артикул','Кратность')


def test_rtp_parser_rejects_whole_article_when_duplicate_has_invalid_pack():
    book = Workbook()
    book.active.title = "Прайс списком"
    book.active.append(["КОД", "Упак"])
    book.active.append(["28200", "15/1"])
    book.active.append(["28200", "100+/1"])
    stream = BytesIO()
    book.save(stream)
    book.close()

    parsed = parse_import_xlsx(stream.getvalue())

    assert "28200" not in parsed.values
    diagnostic = next(
        item for item in parsed.diagnostics
        if item.article == "28200" and item.code == "INVALID_PACK_MULTIPLICITY"
    )
    assert diagnostic.row == 3
    assert diagnostic.message == (
        "Кратность коробки '100+/1' не является точным количеством."
    )
