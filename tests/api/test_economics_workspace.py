from fastapi.testclient import TestClient

import backend.api as api
from backend.main import app
from backend.shipment.store import AnalysisSnapshotStore
from tests.economics.test_workspace import sample_snapshot
from tests.api.test_analysis import _analysis_files, _analysis_data


CLIENT = TestClient(app)
BODY = {'analysis_snapshot_id': 'snap-1', 'target_margin': '0.20',
        'target_roi': '0.40', 'goal': 'margin', 'planned_drr': '0.05'}


def test_report_is_tied_to_current_snapshot_and_has_no_mock_data(monkeypatch):
    store = AnalysisSnapshotStore()
    monkeypatch.setattr(api, 'ANALYSIS_STORE', store)
    snapshot = sample_snapshot()
    snapshot.snapshot_id = 'snap-1'
    store.put(snapshot)
    result = CLIENT.post('/api/economics/workspace', json=BODY)
    assert result.status_code == 200, result.text
    payload = result.json()
    assert payload['snapshot_id'] == 'snap-1'
    assert payload['workspace']['products'][0]['article'] == '26572'
    assert payload['workspace']['modeled_shortfall'] == '130'
    assert payload['workspace']['products'][0]['applied_drr_rate'] == '0'
    assert payload['workspace']['period']['to'] == '2026-09-27'
    store.clear()
    stale = CLIENT.post('/api/economics/workspace', json=BODY)
    assert stale.status_code == 409
    assert stale.json()['error']['code'] == 'ANALYSIS_SNAPSHOT_STALE'


def test_invalid_money_input_fails_closed(monkeypatch):
    store = AnalysisSnapshotStore()
    monkeypatch.setattr(api, 'ANALYSIS_STORE', store)
    snapshot = sample_snapshot()
    snapshot.snapshot_id = 'snap-1'
    store.put(snapshot)
    response = CLIENT.post('/api/economics/workspace', json={**BODY, 'planned_drr': 'NaN'})
    assert response.status_code == 400
    assert response.json()['error']['code'] == 'INVALID_ECONOMICS_SCENARIO'


def test_completed_file_analysis_drives_workspace_without_client_math(monkeypatch):
    monkeypatch.setattr(api, 'ANALYSIS_STORE', AnalysisSnapshotStore())
    analyzed = CLIENT.post('/api/analysis', files=_analysis_files(), data=_analysis_data())
    assert analyzed.status_code == 200, analyzed.text
    snapshot = analyzed.json()['snapshot']
    assert snapshot['economics_settings']['advertising_rate'] == '0.01'
    result = CLIENT.post('/api/economics/workspace', json={
        **BODY, 'analysis_snapshot_id': snapshot['snapshot_id']})
    assert result.status_code == 200, result.text
    product = next(p for p in result.json()['workspace']['products'] if p['sku'] == 'SKU-1')
    assert product['article'] == 'ART-1'
    assert product['real_drr_rate'] is None
    assert 'order_revenue_evidence' not in snapshot
    assert 'daily_order_evidence' not in snapshot
    assert api.ANALYSIS_STORE.latest().order_revenue_evidence is not None
    assert product['groups']['destination'][0]['routes'][0]['origin'] == 'Москва'


def test_workspace_and_export_use_current_background_prices_and_keep_snapshot_private(monkeypatch, tmp_path):
    from io import BytesIO
    from openpyxl import load_workbook
    from tests.economics.test_buyer_prices import priced_snapshot
    from tests.economics.test_daily_series import order
    original = priced_snapshot([order('2026-09-01', 100, None, 2, 'SKU')])
    fresh = priced_snapshot([order('2026-09-01', 100, 60, 2, 'SKU')]).daily_order_evidence
    original.snapshot_id = 'snap-1'
    original.source_snapshot_id = 'source-prices'
    store = AnalysisSnapshotStore(); store.put(original)
    monkeypatch.setattr(api, 'ANALYSIS_STORE', store)
    monkeypatch.setattr(api, 'PROJECT_PATH', tmp_path / 'project.json')
    current = {'evidence': original.daily_order_evidence, 'pending': True}
    calls = []
    def evidence(source_id, immutable):
        calls.append((source_id, immutable))
        return current['evidence'], {'pending': current['pending'], 'failures': 0}
    monkeypatch.setattr(api.ORDER_PRICE_ENRICHMENT, 'evidence', evidence)
    first = CLIENT.post('/api/economics/workspace', json={**BODY, 'planned_drr': '0'}).json()['workspace']
    assert first['prices_pending'] is True
    assert first['products'][0]['buyer_prices']['pending'] is True
    assert first['products'][0]['buyer_prices']['buyer_price_mean'] is None
    current.update(evidence=fresh, pending=False)
    second = CLIENT.post('/api/economics/workspace', json={**BODY, 'planned_drr': '0'}).json()['workspace']
    assert second['prices_pending'] is False
    assert second['products'][0]['buyer_prices']['buyer_price_mean'] == '60'
    assert second['products'][0]['buyer_prices']['target_buyer_price'] == '60'
    downloaded = CLIENT.post('/api/economics/export', json={**BODY, 'planned_drr': '0'})
    assert downloaded.status_code == 200
    sheet = load_workbook(BytesIO(downloaded.content)).active
    columns = {h.value: sheet.cell(2, h.column).value for h in sheet[1]}
    assert columns['Средняя цена клиента, ₽'] == columns['Цена клиента при цели, ₽'] == 60
    assert calls[0] == ('source-prices', original.daily_order_evidence)
    assert original.daily_order_evidence.days[0].buyer_price_mean is None
