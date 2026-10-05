from datetime import date
from decimal import Decimal as D
from io import BytesIO
from types import SimpleNamespace as NS
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import backend.api as api
from backend.main import app
from backend.ozon.vault import CredentialVault
from backend.ozon.contracts import OzonCredentials
from backend.shipment.store import AnalysisSnapshotStore
from backend.economics.buyouts import FinanceProductLine, FinanceExpense, FinanceSnapshot
from tests.economics.test_workspace import sample_snapshot


@pytest.fixture
def setup(monkeypatch, tmp_path):
    from backend.ozon.finance_store import FinanceSnapshotStore
    monkeypatch.setattr(api, 'PROJECT_PATH', tmp_path/'project.json')
    monkeypatch.setattr(api, 'ANALYSIS_STORE', AnalysisSnapshotStore())
    monkeypatch.setattr(api, 'FINANCE_STORE', FinanceSnapshotStore(), raising=False)
    vault = CredentialVault(tmp_path/'vault.json')
    vault.setup(OzonCredentials('test-account', 'test-key'), 'long-password')
    monkeypatch.setattr(api, 'OZON_VAULT', vault)
    snap = sample_snapshot(); snap.snapshot_id = 'analysis-1'
    snap.buyout_cost_inputs = (NS(sku='100', cost=D('200')), NS(sku='200', cost=D('100')))
    api.ANALYSIS_STORE.put(snap)
    f = FinanceSnapshot('finance-1', vault.credential_context_id(), date(2026, 9, 1), date(2026, 9, 30),
        (FinanceProductLine(date(2026, 9, 1), '100', 3, D('3000'), D('1800')),
         FinanceProductLine(date(2026, 9, 2), '200', 1, D('1000'), D('600'))),
        (FinanceExpense(date(2026, 9, 2), 'advertising', 'Реклама', D('100')),),
        '2026-10-02T10:00:00+00:00')
    api.FINANCE_STORE.put(f)
    return TestClient(app), snap, f


BODY = {'analysis_snapshot_id': 'analysis-1', 'finance_snapshot_id': 'finance-1'}


def test_current_uploaded_costs_and_filtered_xlsx_with_store_summary(setup):
    c, snap, _ = setup
    r = c.post('/api/economics/buyouts/workspace', json={**BODY, 'search': '100'})
    assert r.status_code == 200, r.text
    r = r.json()['workspace']
    assert r['products'][0]['cost'] == '200' and r['selected_totals']['profit'] == '1200'
    assert r['totals']['profit_after_known_expenses'] == '1600'
    export = c.post('/api/economics/buyouts/export', json={**BODY, 'search': '100'})
    assert export.status_code == 200, export.text
    book = load_workbook(BytesIO(export.content))
    assert book.sheetnames == ['Выкупы', 'Итог периода', 'Расходы']
    assert book['Выкупы'].max_row == 2 and book['Выкупы']['A2'].value == '100'
    assert 'Общие расходы всего магазина' in [row[0].value for row in book['Итог периода']]


def test_locked_or_replaced_credentials_cannot_read_cached_finance(setup):
    c, _, _ = setup
    api.OZON_VAULT.lock()
    assert c.post('/api/economics/buyouts/workspace', json=BODY).status_code == 423
    api.OZON_VAULT.setup(OzonCredentials('other-account', 'other-key'), 'long-password')
    assert c.post('/api/economics/buyouts/workspace', json=BODY).status_code == 409


def test_stale_analysis_or_invalid_selection_is_rejected(setup):
    c, _, _ = setup
    for extra in ({'search': []}, {'filter': 'bad'}, {'finance_snapshot_id': 'missing'},
                  {'analysis_snapshot_id': 'missing'}):
        r = c.post('/api/economics/buyouts/workspace', json={**BODY, **extra})
        assert r.status_code in (400, 409), r.text
    r = c.post('/api/economics/buyouts/export', json={**BODY, 'search': 'no match'})
    assert r.status_code == 400


def test_sync_rejects_bad_dates_before_network_and_preserves_prior_snapshot(setup):
    c, _, old = setup
    for extra in ({'period_from': '2026-10-01', 'period_to': '2026-09-01'},
                  {'period_from': [], 'period_to': '2026-09-01'},
                  {'period_from': '2024-01-01', 'period_to': '2026-09-01'}):
        r = c.post('/api/economics/buyouts/sync', json=extra)
        assert r.status_code == 400, r.text
    assert api.FINANCE_STORE.get(old.snapshot_id) is old


def test_stream_success_and_failed_refresh_are_atomic(setup, monkeypatch):
    from tests.ozon.test_finance import Client, posting
    c, _, old = setup
    monkeypatch.setattr(api, 'OZON_CLIENT', NS(bind_context=lambda context: Client([
        {'accruals': [posting()], 'last_id': ''}])))
    r = c.post('/api/economics/buyouts/sync', json={'period_from': '2026-09-01', 'period_to': '2026-09-01'})
    assert r.status_code == 200, r.text
    messages = [json.loads(line) for line in r.text.splitlines() if line]
    result = next(m['data'] for m in messages if m['type'] == 'result')
    assert api.FINANCE_STORE.get(result['finance_snapshot_id']).products[0].quantity == 3
    monkeypatch.setattr(api, 'OZON_CLIENT', NS(bind_context=lambda context: Client([
        {'accruals': [posting(total_amount={'amount': 'NaN', 'currency': 'RUB'})], 'last_id': ''}])))
    r = c.post('/api/economics/buyouts/sync', json={'period_from': '2026-09-01', 'period_to': '2026-09-01'})
    messages = [json.loads(line) for line in r.text.splitlines() if line]
    assert messages[-1]['type'] == 'error' and all(m['type'] != 'result' for m in messages)
    assert api.FINANCE_STORE.get(old.snapshot_id) is old


def test_cost_evidence_remains_backend_only(setup):
    from tests.api.test_analysis import _analysis_files, _analysis_data
    c, _, _ = setup
    assert c.put('/api/project/cost-prices/ART-1', json={'cost': '999'}).status_code == 200
    result = c.post('/api/analysis', files=_analysis_files(), data=_analysis_data())
    assert result.status_code == 200, result.text
    snap = api.ANALYSIS_STORE.latest()
    assert snap.buyout_cost_inputs[0].cost == D('100')
    assert 'buyout_cost_inputs' not in result.json()['snapshot']
    assert 'buyout_cost_inputs' not in api.PROJECT_PATH.read_text()


def test_api_analysis_retains_uploaded_old_sku_cost_only_for_finance(setup, monkeypatch):
    from dataclasses import replace
    from backend.ozon.source_store import OzonSourceSnapshotStore
    from tests.api.test_analysis import _api_parity_fixture, _parity_files, _analysis_data, PRODUCT_HEADERS
    from tests.helpers.xlsx_fixtures import make_xlsx
    c, _, _ = setup
    source = replace(_api_parity_fixture(), credential_context_id=api.OZON_VAULT.credential_context_id())
    monkeypatch.setattr(api, 'OZON_SOURCE_STORE', OzonSourceSnapshotStore())
    api.OZON_SOURCE_STORE.put(source)
    products = make_xlsx(headers=PRODUCT_HEADERS, rows=[
        ['SKU-1', 'ART-1', 100, 99, 1000, '10%', 1],
        ['OLD-SKU', 'OLD-ART', 200, 0, 1000, '10%', 1]])
    response = c.post('/api/analysis', files={
        'tariffs_file': _parity_files()['tariffs_file'],
        'product_economics_file': ('products.xlsx', products)}, data=_analysis_data(
            source_mode='api', source_snapshot_id=source.source_snapshot_id))
    assert response.status_code == 200, response.text
    snap = api.ANALYSIS_STORE.latest()
    costs = {p.sku: p.cost for p in snap.buyout_cost_inputs}
    assert costs['OLD-SKU'] == D('200')
    assert all(p.sku != 'OLD-SKU' for p in snap.unit_economics)
    assert all(p.sku != 'OLD-SKU' for p in snap.decision_rows)
    from backend.project import load_project_if_exists
    assert 'OLD-ART' not in load_project_if_exists(api.PROJECT_PATH).cost_prices
    assert 'buyout_cost_inputs' not in response.json()['snapshot']
    finance = FinanceSnapshot('old-finance', api.OZON_VAULT.credential_context_id(),
        date(2026, 9, 1), date(2026, 9, 30),
        (FinanceProductLine(date(2026, 9, 1), 'OLD-SKU', 1, D('1000'), D('600')),), (), 'now')
    api.FINANCE_STORE.put(finance)
    result = c.post('/api/economics/buyouts/workspace', json={
        'analysis_snapshot_id': snap.snapshot_id, 'finance_snapshot_id': finance.snapshot_id})
    assert result.status_code == 200, result.text
    p = result.json()['workspace']['products'][0]
    assert p['sku'] == 'OLD-SKU' and p['article'] == 'OLD-ART'
    assert p['cost'] == '200' and p['profit'] == '400'


def test_account_switch_during_fetch_cannot_commit_or_emit_result(setup, monkeypatch):
    from tests.ozon.test_finance import Client, posting
    c, _, old = setup
    class SwitchingClient(Client):
        def post_json(self, path, body, *, policy, check_cancelled=None):
            response = super().post_json(path, body, policy=policy, check_cancelled=check_cancelled)
            if path.endswith('/by-day'):
                api.OZON_VAULT.setup(OzonCredentials('other-account', 'other-key'), 'long-password')
            return response
    monkeypatch.setattr(api, 'OZON_CLIENT', NS(bind_context=lambda context: SwitchingClient([
        {'accruals': [posting()], 'last_id': ''}])))
    r = c.post('/api/economics/buyouts/sync', json={'period_from': '2026-09-01', 'period_to': '2026-09-01'})
    messages = [json.loads(line) for line in r.text.splitlines() if line]
    assert messages[-1]['type'] == 'error'
    assert all(m['type'] != 'result' for m in messages)
    assert api.FINANCE_STORE.get(old.snapshot_id) is old


def test_formula_like_names_are_literal_and_fee_only_period_can_export(setup):
    c, snap, f = setup
    snap.decision_rows = (NS(sku='100', article='=1+1', product_name='=SUM(1,2)'),)
    export = c.post('/api/economics/buyouts/export', json={**BODY, 'search': '100'})
    sheet = load_workbook(BytesIO(export.content)).active
    assert sheet['B2'].data_type == 's' and sheet['C2'].data_type == 's'
    from dataclasses import replace
    api.FINANCE_STORE.put(replace(f, products=()))
    export = c.post('/api/economics/buyouts/export', json=BODY)
    assert export.status_code == 200
    assert load_workbook(BytesIO(export.content)).active.max_row == 1


def test_slow_finance_read_keeps_the_stream_alive(setup, monkeypatch):
    import time
    import backend.economics.buyout_api as buyout_api
    c, _, finance = setup
    def slow_fetch(*args, progress_callback):
        progress_callback({'current': 0, 'total': 1, 'day': '2026-09-01', 'stage': 'day'})
        time.sleep(1.2)
        progress_callback({'current': 1, 'total': 1, 'day': '2026-09-01', 'stage': 'complete'})
        return finance
    monkeypatch.setattr(buyout_api, 'fetch_finance', slow_fetch)
    r = c.post('/api/economics/buyouts/sync', json={'period_from': '2026-09-01', 'period_to': '2026-09-01'})
    events = [json.loads(line) for line in r.text.splitlines() if line]
    assert any(e['type'] == 'heartbeat' and e['elapsed_seconds'] >= 1 for e in events)
    assert any(e['type'] == 'progress' and e.get('stage') == 'complete'
               and e.get('elapsed_seconds', 0) >= 1 for e in events)
    assert events[-1]['type'] == 'result'


@pytest.mark.parametrize('code, expected', [
    ('OZON_PERMISSION_DENIED', 'финансовым начислениям'),
    ('OZON_RATE_LIMITED', 'лимит запросов'),
    ('OZON_UNAVAILABLE', 'не отвечает'),
])
def test_finance_error_explains_recovery_without_vendor_secrets(setup, monkeypatch, code, expected):
    import backend.economics.buyout_api as buyout_api
    from backend.ozon.client import OzonClientError
    from backend.ozon.contracts import OzonErrorCode
    c, _, old = setup
    def failed(*args, **kwargs):
        raise OzonClientError(OzonErrorCode(code), 'synthetic-key', vendor_message='synthetic-key')
    monkeypatch.setattr(buyout_api, 'fetch_finance', failed)
    r = c.post('/api/economics/buyouts/sync', json={'period_from': '2026-09-01', 'period_to': '2026-09-01'})
    events = [json.loads(line) for line in r.text.splitlines() if line]
    assert events[-1]['error']['code'] == code
    assert expected in events[-1]['error']['message']
    assert 'synthetic-key' not in r.text
    assert api.FINANCE_STORE.get(old.snapshot_id) is old
