from datetime import date
from decimal import Decimal as D
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

import backend.api as api
from backend.main import app
from backend.ozon.vault import CredentialVault
from backend.ozon.contracts import OzonCredentials
from backend.shipment.store import AnalysisSnapshotStore
from backend.ozon.finance_store import FinanceSnapshotStore
from backend.economics.buyouts import FinanceSnapshot, FinanceProductLine, FinanceExpense
from tests.api.test_analysis import _analysis_files, _analysis_data, PRODUCT_HEADERS
from tests.helpers.xlsx_fixtures import make_xlsx

BODY = dict(target_margin='.2', target_roi='.4', goal='margin', planned_drr='.05')


@pytest.mark.parametrize('uploaded_old_unit', [True, False])
@pytest.mark.parametrize('with_finance', [True, False])
def test_removed_sku_orders_are_in_store_total_without_reactivating_plan(data, uploaded_old_unit, with_finance):
    from dataclasses import replace
    from backend.domain.contracts import OrderLifecycle
    from tests.api.test_analysis import _api_parity_fixture, _parity_files
    c, _, _ = data
    base = _api_parity_fixture()
    old = replace(base.orders[0], sku='OLD', article='OLD-ART', quantity=9)
    cancelled = replace(old, quantity=4, lifecycle=OrderLifecycle.CANCELLED)
    source = replace(base, source_snapshot_id='store-history', orders=base.orders + (old, cancelled),
        credential_context_id=api.OZON_VAULT.credential_context_id())
    api.OZON_SOURCE_STORE.put(source)
    rows = [['SKU-1', 'ART-1', 100, 0, 1000, '10%', 1]]
    if uploaded_old_unit:
        rows.append(['OLD', 'OLD-ART', 100, 0, 1000, '10%', 1])
    files = _parity_files()
    result = c.post('/api/analysis', files={
        'tariffs_file': files['tariffs_file'],
        'product_economics_file': ('costs.xlsx', make_xlsx(headers=PRODUCT_HEADERS, rows=rows)),
    }, data=_analysis_data(source_mode='api', source_snapshot_id=source.source_snapshot_id))
    assert result.status_code == 200, result.text
    snapshot = result.json()['snapshot']
    assert {r['sku'] for r in snapshot['decision_rows']} == {'SKU-1'}
    assert {r['sku'] for r in snapshot['shippable_plan']['lines']} <= {'SKU-1'}
    before = api.wire(api.ANALYSIS_STORE.latest())
    body = {**BODY, 'analysis_snapshot_id': snapshot['snapshot_id'],
        'period_from': '2026-07-01', 'period_to': '2026-07-31', 'mode': 'orders'}
    if with_finance:
        finance = FinanceSnapshot('STORE', api.OZON_VAULT.credential_context_id(),
            date(2026, 7, 1), date(2026, 7, 31), (), (
                FinanceExpense(date(2026, 7, 6), 'advertising', 'Ads', D('50')),
                FinanceExpense(date(2026, 7, 6), 'storage', 'Storage', D('100'))), 'now')
        api.FINANCE_STORE.put(finance)
        body['finance_snapshot_id'] = 'STORE'
    response = c.post('/api/economics/period/workspace', json=body)
    assert response.status_code == 200, response.text
    report = response.json()['workspace']
    assert report['quantity_complete'] is True
    assert report['totals']['qty'] == 35
    old_row = next(p for p in report['products'] if p['sku'] == 'OLD')
    assert old_row['qty'] == 9
    if uploaded_old_unit:
        assert D(old_row['profit']) == D('6120')
        assert D(report['totals']['profit_before_common']) == D('23800')
        if with_finance:
            assert D(report['totals']['profit_after_common']) == D('23650')
            assert report['totals']['partial'] is False
    else:
        assert old_row['profit'] is None
        assert report['totals']['uncovered_known_qty'] == 9
        assert D(report['totals']['profit_before_common']) == D('17680')
        assert report['totals']['partial'] is True
    public = c.post('/api/economics/workspace', json=body)
    assert public.status_code == 200, public.text
    assert {p['sku'] for p in public.json()['workspace']['products']} == {'SKU-1'}
    exported = c.post('/api/economics/period/export', json={**body, 'search': 'OLD'})
    assert exported.status_code == 200, exported.text
    sheet = load_workbook(BytesIO(exported.content))['Товары']
    assert sheet.max_row == 2
    assert 'OLD' in [v.value for v in sheet[2]]
    assert api.wire(api.ANALYSIS_STORE.latest()) == before


@pytest.mark.parametrize('mode', ['orders', 'buyouts'])
def test_historical_quantity_only_sku_remains_filterable_and_exportable(data, mode):
    from dataclasses import replace
    from backend.domain.contracts import OrderLifecycle
    from tests.api.test_analysis import _api_parity_fixture, _parity_files
    c, _, _ = data
    base = _api_parity_fixture()
    old = replace(base.orders[0], sku='OLD', article='OLD-ART', quantity=9,
        lifecycle=OrderLifecycle.IN_PROGRESS)
    source = replace(base, source_snapshot_id='quantity-only-history',
        orders=base.orders + ((old,) if mode == 'orders' else ()),
        credential_context_id=api.OZON_VAULT.credential_context_id())
    api.OZON_SOURCE_STORE.put(source)
    files = _parity_files()
    analysis = c.post('/api/analysis', files={
        'tariffs_file': files['tariffs_file'],
        'product_economics_file': ('costs.xlsx', make_xlsx(headers=PRODUCT_HEADERS,
            rows=[['SKU-1', 'ART-1', 100, 0, 1000, '10%', 1]])),
    }, data=_analysis_data(source_mode='api', source_snapshot_id=source.source_snapshot_id))
    assert analysis.status_code == 200, analysis.text
    snapshot = analysis.json()['snapshot']
    assert {r['sku'] for r in snapshot['decision_rows']} == {'SKU-1'}
    before = api.wire(api.ANALYSIS_STORE.latest())
    finance = FinanceSnapshot('QUANTITY', api.OZON_VAULT.credential_context_id(),
        date(2026, 7, 1), date(2026, 7, 31),
        (FinanceProductLine(date(2026, 7, 6), 'OLD', 9, D('9000'), D('1')),) if mode == 'buyouts' else (),
        (), 'now')
    api.FINANCE_STORE.put(finance)
    body = {**BODY, 'analysis_snapshot_id': snapshot['snapshot_id'],
        'period_from': '2026-07-01', 'period_to': '2026-07-31',
        'mode': mode, 'finance_snapshot_id': 'QUANTITY'}
    total = c.post('/api/economics/period/workspace', json=body)
    assert total.status_code == 200, total.text
    report = total.json()['workspace']
    assert report['totals']['qty'] == (35 if mode == 'orders' else 9)
    assert report['totals']['uncovered_known_qty'] == 9
    for selection in ({'search': 'OLD'}, {'filter': 'incomplete'}):
        response = c.post('/api/economics/period/workspace', json={**body, **selection})
        assert response.status_code == 200, response.text
        products = response.json()['workspace']['products']
        assert [p['sku'] for p in products] == ['OLD']
        assert products[0]['qty'] == 9 and products[0]['profit'] is None
        assert products[0]['partial'] is True
        exported = c.post('/api/economics/period/export', json={**body, **selection})
        assert exported.status_code == 200, exported.text
        sheet = load_workbook(BytesIO(exported.content))['Товары']
        assert sheet.max_row == 2 and 'OLD' in [v.value for v in sheet[2]]
    public = c.post('/api/economics/workspace', json=body)
    assert public.status_code == 200, public.text
    assert {p['sku'] for p in public.json()['workspace']['products']} == {'SKU-1'}
    assert api.wire(api.ANALYSIS_STORE.latest()) == before


def test_pricing_does_not_block_other_async_requests(data, monkeypatch):
    import asyncio
    from threading import Event
    import httpx
    from backend.economics import workspace_pricing
    from backend.security import LOCAL_SESSION_HEADER, current_local_session_token
    _, sid, period = data
    entered, release = Event(), Event()
    original = workspace_pricing.enrich_pricing
    def slow(*args, **kwargs):
        entered.set()
        assert release.wait(timeout=1), 'pricing blocked the application event loop'
        return original(*args, **kwargs)
    monkeypatch.setattr(workspace_pricing, 'enrich_pricing', slow)
    async def run():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),
            base_url='http://127.0.0.1',headers={LOCAL_SESSION_HEADER:current_local_session_token()}) as client:
            request = asyncio.create_task(client.post('/api/economics/workspace',json={**BODY,**period,
                'analysis_snapshot_id':sid}))
            try:
                assert await asyncio.to_thread(entered.wait, 2)
                health = await asyncio.wait_for(client.get('/api/project/mappings'),timeout=.5)
                assert health.status_code == 200
            finally:
                release.set()
            result = await request
            assert result.status_code == 200, result.text
    asyncio.run(run())


@pytest.fixture
def data(monkeypatch, tmp_path):
    monkeypatch.setattr(api, 'PROJECT_PATH', tmp_path / 'project.json')
    monkeypatch.setattr(api, 'ANALYSIS_STORE', AnalysisSnapshotStore())
    monkeypatch.setattr(api, 'FINANCE_STORE', FinanceSnapshotStore())
    vault = CredentialVault(tmp_path / 'vault.json')
    vault.setup(OzonCredentials('account', 'key'), 'long-password')
    monkeypatch.setattr(api, 'OZON_VAULT', vault)
    c = TestClient(app)
    f = _analysis_files()
    f['product_economics_file'] = ('products.xlsx', make_xlsx(headers=PRODUCT_HEADERS,
        rows=[['SKU-1', 'ART-1', 100, 20, 1000, '10%', 1],
              ['NEW', 'NEW-ART', 100, 0, 1000, '10%', 1]]))
    result = c.post('/api/analysis', files=f, data=_analysis_data())
    assert result.status_code == 200, result.text
    sid = result.json()['snapshot']['snapshot_id']
    period = {'period_from':'2026-08-20', 'period_to':'2026-08-20'}
    finance = FinanceSnapshot('FIN', vault.credential_context_id(), date(2026, 8, 20), date(2026, 8, 20),
        (FinanceProductLine(date(2026, 8, 20), 'SKU-1', 2, D('1'), D('7')),),
        (FinanceExpense(date(2026, 8, 20), 'advertising', 'Ads', D('50'), 'NEW'),), 'now')
    api.FINANCE_STORE.put(finance)
    return c, sid, period


def test_unit_basis_is_private_and_target_works_for_no_sales(data):
    c, sid, period = data
    r = c.post('/api/economics/workspace', json={**BODY, **period, 'analysis_snapshot_id':sid,
        'scenario_routes': {'NEW': [['Москва', 'Москва']]}})
    assert r.status_code == 200, r.text
    new = next(p for p in r.json()['workspace']['products'] if p['sku'] == 'NEW')
    assert new['calculation_kind'] == 'no_sales'
    assert new['target_price_all_routes'] is not None
    assert new['profit_per_unit_before_ads'] is not None
    wire = api.wire(api.ANALYSIS_STORE.latest())
    for field in ('economics_pricing_inputs', 'economics_tariffs', 'economics_order_quantities',
                  'economics_order_quantities_complete', 'economics_catalog_identities'):
        assert field not in wire
        assert field not in api.PROJECT_PATH.read_text()


def test_modes_use_same_basis_and_expenses_and_export(data):
    c, sid, period = data
    body = {**BODY, **period, 'analysis_snapshot_id':sid, 'finance_snapshot_id':'FIN'}
    reports = []
    for mode in ('orders', 'buyouts'):
        r = c.post('/api/economics/period/workspace', json={**body, 'mode':mode})
        assert r.status_code == 200, r.text
        reports.append(r.json()['workspace'])
    a, b = reports
    assert a['pricing_basis_id'] == b['pricing_basis_id']
    assert a['totals']['advertising_total'] == b['totals']['advertising_total'] == '50'
    assert D(b['totals']['profit_before_common']) == D(a['totals']['profit_before_common']) * 2
    r = c.post('/api/economics/period/export', json={**body, 'mode':'buyouts', 'search':'ART-1'})
    assert r.status_code == 200, r.text
    book = load_workbook(BytesIO(r.content))
    assert book.sheetnames == ['Итог', 'Товары', 'Расходы']
    assert book['Товары'].max_row == 2


@pytest.mark.parametrize('extra,complete,qty', [
    ('SKU-1;2;1000;;Москва;В пути;2026-08-20T10:00:00\n', True, 3),
    ('SKU-1;bad;1000;Москва;Москва;В пути;2026-08-20T10:00:00\n', False, None),
    ('SKU-1;2;1000;Москва;Москва;В пути;bad-date\n', False, None),
    ('SKU-1;2;1000;Москва;Москва;unknown;2026-08-20T10:00:00\n', False, None),
])
def test_files_order_quantity_preserves_originless_orders_and_reports_gaps(data, extra, complete, qty):
    c, _, period = data
    files = _analysis_files()
    name, rows = files['orders_file']
    files['orders_file'] = (name, rows + extra.encode())
    analysis = c.post('/api/analysis', files=files, data=_analysis_data())
    assert analysis.status_code == 200, analysis.text
    sid = analysis.json()['snapshot']['snapshot_id']
    result = c.post('/api/economics/period/workspace', json={**BODY, **period,
        'analysis_snapshot_id':sid, 'mode':'orders'})
    assert result.status_code == 200, result.text
    report = result.json()['workspace']
    assert report['quantity_complete'] is complete
    assert report['totals']['qty'] == qty
    if complete:
        assert D(report['totals']['profit_before_common']) == D('2040')
        assert sum(r.observed_qty for r in api.ANALYSIS_STORE.latest().route_economics) == 1
    else:
        assert report['totals']['partial'] is True


def test_expense_load_optional_for_orders_and_credentials_guard_cache(data):
    c, sid, period = data
    body = {**BODY, **period, 'analysis_snapshot_id':sid, 'mode':'orders'}
    r = c.post('/api/economics/period/workspace', json=body)
    assert r.status_code == 200, r.text
    assert r.json()['workspace']['totals']['profit_after_common'] is None
    api.OZON_VAULT.lock()
    assert c.post('/api/economics/period/workspace', json={**body,'finance_snapshot_id':'FIN'}).status_code == 423
    api.OZON_VAULT.setup(OzonCredentials('other', 'key2'), 'long-password')
    assert c.post('/api/economics/period/workspace', json={**body,'finance_snapshot_id':'FIN'}).status_code == 409


def test_new_upload_cost_overrides_older_manual_cost(data):
    c, sid, period = data
    c.put('/api/project/cost-prices/ART-1', json={'cost':'999'})
    r = c.post('/api/economics/workspace', json={**BODY, **period, 'analysis_snapshot_id':sid})
    assert r.status_code == 200, r.text
    assert next(p for p in r.json()['workspace']['products'] if p['sku'] == 'SKU-1')['cost'] == '100'


def test_cluster_math_uses_uploaded_cost_and_canonical_unit(data):
    c, sid, period = data
    c.put('/api/project/cost-prices/ART-1', json={'cost':'999'})
    r = c.post('/api/economics/workspace', json={**BODY, **period, 'analysis_snapshot_id':sid})
    p = next(p for p in r.json()['workspace']['products'] if p['sku'] == 'SKU-1')
    route = p['groups']['destination'][0]['routes'][0]
    assert route['cost'] == '100'
    assert D(route['profit']) == D(p['profit_per_unit'])
    assert D(route['profit_before_ads']) == D(p['profit_per_unit_before_ads'])
    assert D(p['line_items']['PROFIT_PER_UNIT']) == D(p['profit_per_unit'])
    assert p['basis_period'] == r.json()['workspace']['period']
    before = api.wire(api.ANALYSIS_STORE.latest())
    c.post('/api/economics/period/workspace', json={**BODY, **period,'analysis_snapshot_id':sid, 'mode':'buyouts','finance_snapshot_id':'FIN'})
    assert api.wire(api.ANALYSIS_STORE.latest()) == before


def test_explicit_current_cost_changes_unit_basis_without_replacing_upload(data):
    c, sid, period = data
    body = {**BODY, **period, 'analysis_snapshot_id':sid}
    first = c.post('/api/economics/workspace', json=body).json()['workspace']
    changed = c.post('/api/economics/workspace', json={**body,'per_sku_cost':{'SKU-1':'120'}}).json()['workspace']
    p = next(p for p in changed['products'] if p['sku'] == 'SKU-1')
    original = next(p for p in first['products'] if p['sku'] == 'SKU-1')
    assert p['cost'] == '120'
    assert p['cost_source'] == 'manual_current'
    assert D(p['profit_per_unit_before_ads']) == D(original['profit_per_unit_before_ads'])-20
    assert next(p for p in c.post('/api/economics/workspace', json=body).json()['workspace']['products'] if p['sku']=='SKU-1')['cost'] == '100'
    assert c.post('/api/economics/workspace', json={**body,'per_sku_cost':{'SKU-1':'NaN'}}).status_code == 400


def test_unit_export_preserves_original_columns_and_explains_sources(data):
    c, sid, period = data
    body = {**BODY, **period,'analysis_snapshot_id':sid,'search':'NEW', 'use_planned_drr':True,
            'scenario_routes':{'NEW':[['Москва','Москва']]}}
    response = c.post('/api/economics/export',json=body)
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    columns = {h.value: sheet.cell(2,h.column).value for h in sheet[1]}
    assert [h.value for h in sheet[1]][:3] == ['Артикул','Товар','Текущая цена, ₽']
    assert columns['SKU'] == 'NEW'
    assert columns['Источник ДРР в расчёте'] == 'План · отчёт отсутствует'
    assert columns['Режим юнитки'] == 'Без продаж'
    assert columns['Прибыль до рекламы, ₽ / шт.'] > columns['Прибыль, ₽ / шт.']
    assert columns['Расчётные маршруты'] == 'Москва → Москва'


def test_old_buyout_sku_with_complete_inputs_is_counted_without_catalog_reactivation(data):
    from dataclasses import replace
    from backend.domain.contracts import SourceMode
    c, sid, period = data
    snapshot = api.ANALYSIS_STORE.latest()
    snapshot = replace(snapshot, source_mode=SourceMode.API, economics_catalog_identities=tuple(
        x for x in snapshot.economics_catalog_identities if x[0]!='NEW'))
    api.ANALYSIS_STORE.put(snapshot)
    old_finance = api.FINANCE_STORE.get('FIN')
    api.FINANCE_STORE.put(replace(old_finance,products=old_finance.products + (
        FinanceProductLine(date(2026,8,20),'NEW',2,D('1'),D('7')),)))
    body = {**BODY,**period,'analysis_snapshot_id':sid,'scenario_routes':{'NEW':[['Москва','Москва']]}}
    unit = c.post('/api/economics/workspace',json=body).json()['workspace']
    assert {p['sku'] for p in unit['products']} == {'SKU-1'}
    plan = api.wire(snapshot.decision_rows)
    response = c.post('/api/economics/period/workspace',json={**body,'mode':'buyouts','finance_snapshot_id':'FIN'})
    assert response.status_code == 200,response.text
    old = next(p for p in response.json()['workspace']['products'] if p['sku']=='NEW')
    assert old['qty'] == 2 and old['profit'] is not None
    assert api.wire(snapshot.decision_rows) == plan


def test_api_catalog_card_without_orders_or_clusters_can_price_from_cost_file():
    from dataclasses import replace
    from tests.api.test_analysis import _api_parity_fixture, _parity_files
    from backend.ozon.adapters.product_facts import ProductApiFacts
    c = TestClient(app)
    base = _api_parity_fixture()
    source = replace(base,source_snapshot_id='new-pricing',orders=(),availability=(),
        operational_seller_stock=(),product_facts=(ProductApiFacts('NEW','NEW-ART',77),))
    api.OZON_SOURCE_STORE.put(source)
    files = _parity_files()
    response = c.post('/api/analysis',files={'tariffs_file':files['tariffs_file'],
        'product_economics_file':('cost.xlsx',make_xlsx(headers=PRODUCT_HEADERS,
            rows=[['NEW','NEW-ART',100,0,1000,'10%',1]]))},
        data=_analysis_data(source_mode='api',source_snapshot_id=source.source_snapshot_id))
    assert response.status_code == 200,response.text
    snapshot = response.json()['snapshot']
    assert snapshot['decision_rows'] == []
    body = {**BODY,'analysis_snapshot_id':snapshot['snapshot_id'],
        'scenario_routes':{'NEW':[['Москва','Москва']]}}
    unit = c.post('/api/economics/workspace',json=body)
    assert unit.status_code == 200,unit.text
    p = unit.json()['workspace']['products'][0]
    assert p['sku']=='NEW' and p['calculation_kind']=='no_sales'
    assert p['target_price_all_routes'] is not None


def test_configured_new_sku_is_complete_even_with_zero_sales(data):
    c,sid,period=data
    response=c.post('/api/economics/workspace',json={**BODY,**period,'analysis_snapshot_id':sid,
        'scenario_routes':{'NEW':[['Москва','Москва']]},'filter':'incomplete'})
    assert response.status_code == 200,response.text
    assert response.json()['workspace']['products'] == []
