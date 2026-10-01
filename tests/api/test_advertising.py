from decimal import Decimal
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

import backend.api as api
from backend.main import app
from backend.shipment.store import AnalysisSnapshotStore
from tests.api.test_economics_workspace import BODY
from tests.economics.test_workspace import sample_snapshot
from tests.economics.test_advertising import advertising_xlsx, evidence, order


@pytest.fixture
def client(monkeypatch):
    snapshot = sample_snapshot()
    snapshot.snapshot_id = 'snap-1'
    snapshot.order_revenue_evidence = evidence([order('2026-09-01', 100, sku='SKU'),
                                               order('2026-09-02', 300, sku='SKU')])
    store = AnalysisSnapshotStore()
    store.put(snapshot)
    monkeypatch.setattr(api, 'ANALYSIS_STORE', store)
    return TestClient(app)


def test_missing_report_is_na_and_does_not_use_previous_manual_model_drr(client):
    report = client.post('/api/economics/workspace', json={**BODY, 'modeled_drr': '.5'}).json()['workspace']
    product = report['products'][0]
    assert product['real_drr_rate'] is None
    assert product['margin'] is not None and product['roi'] is not None
    assert product['applied_drr_rate'] == '0' and product['drr_zero_assumed'] is True
    assert product['target_price_all_routes'] is not None


def test_batch_import_is_per_file_and_persists_real_model_and_export(client):
    report = advertising_xlsx(rows=[['01.09.2026', 'SKU', 'Товар', 10],
                                    ['02.09.2026', 'SKU', 'Товар', 30]])
    result = client.post('/api/economics/advertising/import',
                         data={'analysis_snapshot_id': 'snap-1'},
                         files=[('files', ('report.xlsx', report)), ('files', ('bad.xlsx', b'not xlsx'))])
    assert result.status_code == 200, result.text
    assert [r['status'] for r in result.json()['files']] == ['imported', 'error']
    assert result.json()['files'][0]['matched_products'] == [
        {'sku':'SKU','article':'26572','name':'Розетка'}]
    assert api.PROJECT_PATH.with_name('advertising.json').exists()
    product = client.post('/api/economics/workspace', json=BODY).json()['workspace']['products'][0]
    assert product['real_drr_rate'] == '0.1'
    assert product['profit_per_unit'] == '1.333333333333333333333333333333333333333'
    export = client.post('/api/economics/export', json=BODY)
    assert export.status_code == 200
    sheet = load_workbook(BytesIO(export.content), data_only=True).active
    headers = [cell.value for cell in sheet[1]]
    assert 'Реальный ДРР, %' in headers
    assert sheet.cell(2, headers.index('Реальный ДРР, %') + 1).value == .1


def test_stale_import_cannot_change_saved_advertising(client):
    result = client.post('/api/economics/advertising/import',
                         data={'analysis_snapshot_id': 'old'},
                         files={'files': ('report.xlsx', advertising_xlsx())})
    assert result.status_code == 409
    assert not api.PROJECT_PATH.with_name('advertising.json').exists()


def upload(client, content=None, campaign='123'):
    content = content or advertising_xlsx(campaign=campaign, rows=[['01.09.2026', 'SKU', 'Товар', 40]])
    return client.post('/api/economics/advertising/import', data={'analysis_snapshot_id':'snap-1'},
                       files={'files':('report.xlsx',content)})


def test_duplicate_correction_multiple_campaigns_and_delete(client):
    assert upload(client).status_code == 200
    path = api.PROJECT_PATH.with_name('advertising.json')
    previous = path.read_bytes()
    assert upload(client).json()['files'][0]['status'] == 'duplicate'
    assert path.read_bytes() == previous
    corrected = advertising_xlsx(rows=[['01.09.2026','SKU','Товар',20]])
    assert upload(client, corrected).json()['changed'] is True
    assert upload(client, campaign='456').json()['changed'] is True
    product = client.post('/api/economics/workspace', json=BODY).json()['workspace']['products'][0]
    assert product['real_drr_rate'] == '0.15'  # (20 + 40) / (100 + 300)
    removed = client.request('DELETE','/api/economics/advertising/456',json={'analysis_snapshot_id':'snap-1'})
    assert removed.status_code == 200
    assert len(removed.json()['campaigns']) == 1
    assert client.request('DELETE','/api/economics/advertising/123',json={'analysis_snapshot_id':'snap-1'}).status_code == 200
    product = client.post('/api/economics/workspace', json=BODY).json()['workspace']['products'][0]
    assert product['real_drr_rate'] is None
    assert upload(client).json()['files'][0]['status'] == 'imported'


def test_unmatched_sku_is_visible_and_does_not_create_an_expense_directory(client):
    result = upload(client, advertising_xlsx())
    assert result.json()['changed'] is False
    assert result.json()['files'][0]['unmatched_skus'] == ['100']
    assert result.json()['files'][0]['status'] == 'unmatched'
    assert not api.PROJECT_PATH.with_name('advertising.json').exists()


def test_failed_atomic_write_preserves_last_report_and_can_retry(client, monkeypatch):
    assert upload(client).status_code == 200
    path = api.PROJECT_PATH.with_name('advertising.json')
    before = path.read_bytes()
    save = api.save_advertising
    def fail(*_): raise OSError('disk unavailable')
    monkeypatch.setattr(api, 'save_advertising', fail)
    corrected = advertising_xlsx(rows=[['01.09.2026','SKU','Товар',20]])
    assert upload(client, corrected).status_code == 500
    assert path.read_bytes() == before
    monkeypatch.setattr(api, 'save_advertising', save)
    assert upload(client, corrected).status_code == 200


def test_snapshot_change_during_parser_rejects_import(client, monkeypatch):
    parse = api.import_advertising
    def change_snapshot(*args):
        result = parse(*args)
        api.ANALYSIS_STORE.clear()
        return result
    monkeypatch.setattr(api, 'import_advertising', change_snapshot)
    assert upload(client).status_code == 409
    assert not api.PROJECT_PATH.with_name('advertising.json').exists()


def test_expense_change_during_export_rejects_stale_download(client, monkeypatch):
    assert upload(client).status_code == 200
    export = api.export_economics
    def change_expense(report):
        result = export(report)
        assert upload(client, advertising_xlsx(rows=[['01.09.2026','SKU','Товар',20]])).status_code == 200
        return result
    monkeypatch.setattr(api, 'export_economics', change_expense)
    response = client.post('/api/economics/export', json=BODY)
    assert response.status_code == 409
    assert response.json()['error']['code'] == 'ECONOMICS_INPUT_CHANGED'


def test_real_file_analysis_supplies_all_order_revenue_and_preserves_plan(client):
    from tests.api.test_analysis import _analysis_files, _analysis_data
    files = _analysis_files()
    files['orders_file'] = ('orders.csv', ('SKU;Количество;Цена продавца;Кластер отгрузки;Кластер доставки;Статус;Принят в обработку\n'
        'SKU-1;1;1000;Москва;Москва;Доставлен;2026-09-01T10:00:00\n'
        'SKU-1;2;1500;Москва;Москва;Отменён;2026-09-02T10:00:00\n').encode())
    analyzed = client.post('/api/analysis',files=files,data=_analysis_data(as_of='2026-09-30',
        orders_period_from='2026-08-06',orders_period_to='2026-09-30'))
    assert analyzed.status_code == 200, analyzed.text
    snapshot = analyzed.json()['snapshot']
    immutable = api.wire(api.ANALYSIS_STORE.latest())
    assert 'order_revenue_evidence' not in snapshot
    result = client.post('/api/economics/advertising/import',data={'analysis_snapshot_id':snapshot['snapshot_id']},
        files={'files':('ads.xlsx',advertising_xlsx(rows=[['01.09.2026','SKU-1','Товар',100],['02.09.2026','SKU-1','Товар',300]]))})
    assert result.status_code == 200, result.text
    report = client.post('/api/economics/workspace',json={**BODY,'analysis_snapshot_id':snapshot['snapshot_id']}).json()['workspace']
    product = next(p for p in report['products'] if p['sku']=='SKU-1')
    assert product['real_drr_rate'] == '0.1'
    assert Decimal(product['advertising']['order_revenue']) == 4000
    assert api.wire(api.ANALYSIS_STORE.latest()) == immutable
