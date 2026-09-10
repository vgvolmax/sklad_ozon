"""API trust-boundary and real XLSX/ZIP byte regressions."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

import backend.api as api_module
from backend.ozon.draft_contracts import ValidationState
from backend.shipment.contracts import ShipmentPlan
from backend.shipment.ranking import rank_outcomes
from tests.api.test_analysis import CLIENT
from tests.shipment.test_ranking import assignment, outcome, scenario


def _stored_plan(*, clusters=("M",), blank_article=False, available=True):
    original=outcome("cs_export",ValidationState.ACCEPTED if available else ValidationState.REJECTED,
                     accepted=available)
    if len(clusters)>1 or blank_article:
        rows=tuple(replace(assignment(cluster,6),article="" if blank_article else "40750") for cluster in clusters)
        candidate=replace(original.candidate,cluster_ids=clusters,assignments=rows,
                          total_qty=6*len(rows),total_volume_l=sum((row.total_volume_l for row in rows),start=rows[0].total_volume_l*0))
        validation=replace(original.validation,accepted_assignments=rows if available else ())
        original=replace(original,candidate=candidate,validation=validation,
                         unresolved_assignments=() if available else rows)
    ranked,unavailable=rank_outcomes((original,),scenario())
    plan=ShipmentPlan("sp_export",None,"as_1","spp_1",scenario().date_from,"ss_export",ranked,unavailable,())
    api_module.SHIPMENT_PLAN_STORE.clear(); api_module.SHIPMENT_PLAN_STORE.put(plan)
    return plan


def test_export_single_cluster_returns_valid_exact_xlsx():
    _stored_plan()
    response=CLIENT.post("/api/shipment/export",json={"shipment_plan_id":"sp_export","option_id":"cs_export"})
    assert response.status_code==200,response.text
    assert response.headers["content-type"]=="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["content-disposition"].endswith('.xlsx"')
    sheet=load_workbook(BytesIO(response.content)).active
    assert tuple(sheet.cell(1,index).value for index in range(1,4))==("артикул","имя (необязательно)","количество")


def test_export_multiple_clusters_returns_only_one_xlsx_per_cluster():
    _stored_plan(clusters=("M","K"))
    response=CLIENT.post("/api/shipment/export",json={"shipment_plan_id":"sp_export","option_id":"cs_export"})
    assert response.status_code==200,response.text
    assert response.headers["content-type"]=="application/zip"
    with ZipFile(BytesIO(response.content)) as archive:
        assert len(archive.namelist())==2
        assert all(name.endswith(".xlsx") and "/" not in name for name in archive.namelist())


@pytest.mark.parametrize("field",["quantity","rows","article","sku","clusters","timeslot","accepted_assignments"])
def test_export_rejects_browser_evidence(field):
    response=CLIENT.post("/api/shipment/export",json={"shipment_plan_id":"sp_x","option_id":"cs_x",field:999})
    assert response.status_code==400
    assert response.json()["error"]["code"]=="UNSUPPORTED_FIELD"


def test_export_unknown_plan_is_not_reconstructed():
    api_module.SHIPMENT_PLAN_STORE.clear()
    response=CLIENT.post("/api/shipment/export",json={"shipment_plan_id":"sp_missing","option_id":"cs_x"})
    assert response.status_code==404 and response.json()["error"]["code"]=="SHIPMENT_PLAN_NOT_FOUND"


def test_export_unknown_and_unavailable_options():
    _stored_plan(available=False)
    unknown=CLIENT.post("/api/shipment/export",json={"shipment_plan_id":"sp_export","option_id":"cs_missing"})
    unavailable=CLIENT.post("/api/shipment/export",json={"shipment_plan_id":"sp_export","option_id":"cs_export"})
    assert unknown.status_code==404 and unknown.json()["error"]["code"]=="SHIPMENT_OPTION_NOT_FOUND"
    assert unavailable.status_code==409 and unavailable.json()["error"]["code"]=="EXPORT_OPTION_NOT_EXPORTABLE"


def test_export_preserves_renderer_error_code():
    _stored_plan(blank_article=True)
    response=CLIENT.post("/api/shipment/export",json={"shipment_plan_id":"sp_export","option_id":"cs_export"})
    assert response.status_code==409
    assert response.json()["error"]["code"]=="EXPORT_ARTICLE_REQUIRED"
