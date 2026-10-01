from io import BytesIO
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import backend.api as api
from backend.main import app
from backend.project import load_project
from backend.shipment.store import AnalysisSnapshotStore
from tests.api.test_analysis import _analysis_data, _analysis_files
from tests.api.test_economics_workspace import BODY
from tests.economics.test_workspace import sample_snapshot


client = TestClient(app)


def test_cost_edit_survives_restart_and_reimport(tmp_path, monkeypatch):
    path = tmp_path / 'project.json'
    monkeypatch.setattr(api, 'PROJECT_PATH', path)
    response = client.post('/api/analysis', files=_analysis_files(), data=_analysis_data())
    assert response.status_code == 200
    items = client.get('/api/project/cost-prices').json()['items']
    assert items[0]['cost'] == '100'
    saved = client.put('/api/project/cost-prices/ART-1', json={'cost': '125.50'})
    assert saved.status_code == 200, saved.text
    assert saved.json()['item']['source'] == 'manual'
    assert load_project(path).cost_prices['ART-1'].cost == Decimal('125.50')
    api.ANALYSIS_STORE.clear()
    response = client.post('/api/analysis', files=_analysis_files(), data=_analysis_data())
    assert response.status_code == 200, response.text
    assert all(x['cost'] == '125.5' for x in response.json()['economics'])
    exported = client.get('/api/project/cost-prices/export')
    assert exported.status_code == 200
    book = load_workbook(BytesIO(exported.content))
    sheet = book.active
    assert sheet.cell(2, 1).value == 'ART-1'
    assert sheet.cell(2, 3).value == 125.5
    assert sheet.cell(2, 3).data_type == 'n'
    assert sheet.freeze_panes == 'A2' and sheet.auto_filter.ref


@pytest.mark.parametrize('cost', ['NaN', '-1', '', True, 'Infinity', None])
def test_invalid_cost_never_overwrites_saved_value(tmp_path, monkeypatch, cost):
    path = tmp_path / 'project.json'
    monkeypatch.setattr(api, 'PROJECT_PATH', path)
    assert client.put('/api/project/cost-prices/A', json={'cost': '40'}).status_code == 200
    previous = path.read_bytes()
    result = client.put('/api/project/cost-prices/A', json={'cost': cost})
    assert result.status_code == 400
    assert path.read_bytes() == previous


def test_economics_updates_cost_without_mutating_snapshot(tmp_path, monkeypatch):
    monkeypatch.setattr(api, 'PROJECT_PATH', tmp_path / 'project.json')
    store = AnalysisSnapshotStore()
    monkeypatch.setattr(api, 'ANALYSIS_STORE', store)
    snapshot = sample_snapshot()
    snapshot.snapshot_id = 'snap-1'
    store.put(snapshot)
    assert client.put('/api/project/cost-prices/26572', json={'cost': '50'}).status_code == 200
    result = client.post('/api/economics/workspace', json={**BODY, 'modeled_drr': '.10'})
    assert result.status_code == 200, result.text
    product = result.json()['workspace']['products'][0]
    assert product['cost'] == '50' and product['cost_source'] == 'manual'
    assert product['assumed_drr_rate'] == '0.1'
    assert product['profit_per_unit'].startswith('-8.666')
    assert snapshot.unit_economics[0].cost == Decimal('40')
    assert snapshot.route_economics[0].current_profit_per_unit == Decimal('13')


def test_economics_excel_has_one_product_row_and_numeric_rates(tmp_path, monkeypatch):
    monkeypatch.setattr(api, 'PROJECT_PATH', tmp_path / 'project.json')
    store = AnalysisSnapshotStore()
    monkeypatch.setattr(api, 'ANALYSIS_STORE', store)
    snapshot = sample_snapshot()
    snapshot.snapshot_id = 'snap-1'
    store.put(snapshot)
    result = client.post('/api/economics/export', json={**BODY, 'modeled_drr': '.10'})
    assert result.status_code == 200, result.text
    sheet = load_workbook(BytesIO(result.content)).active
    assert sheet.max_row == 2
    assert [c.value for c in sheet[1]] == ['Артикул', 'Товар', 'Текущая цена, ₽',
        'ДРР по плану, %', 'Маржа, %', 'ROI, %', 'Плановая маржа, %', 'Необходимая цена, ₽']
    assert sheet.cell(2, 1).value == '26572'
    assert sheet.cell(2, 3).value == 100
    assert sheet.cell(2, 4).value == .05
    assert sheet.cell(2, 7).value == .20
    assert sheet.cell(2, 4).number_format == '0.0%'
    assert sheet.freeze_panes == 'C2' and sheet.auto_filter.ref == 'A1:H2'
    api.ANALYSIS_STORE.clear()
    assert client.post('/api/economics/export', json=BODY).status_code == 409


def test_old_imported_cost_cannot_replace_conflicting_current_sku_costs(tmp_path, monkeypatch):
    from backend.project import CostPriceRecord, Project, save_project_atomic
    path = tmp_path / 'project.json'
    monkeypatch.setattr(api, 'PROJECT_PATH', path)
    save_project_atomic(path, Project(cost_prices={
        '26572': CostPriceRecord(Decimal('30'), 'import', 'now')}))
    store = AnalysisSnapshotStore()
    monkeypatch.setattr(api, 'ANALYSIS_STORE', store)
    snapshot = sample_snapshot()
    snapshot.snapshot_id = 'snap-1'
    store.put(snapshot)
    response = client.post('/api/economics/workspace', json=BODY)
    assert response.status_code == 200
    assert response.json()['workspace']['products'][0]['cost'] == '40'


def test_duplicate_article_is_a_blocking_excel_conflict(tmp_path, monkeypatch):
    from copy import deepcopy
    monkeypatch.setattr(api, 'PROJECT_PATH', tmp_path / 'project.json')
    snapshot = sample_snapshot()
    snapshot.snapshot_id = 'snap-1'
    other_unit = deepcopy(snapshot.unit_economics[0])
    other_unit.sku = 'OTHER'
    snapshot.unit_economics += (other_unit,)
    other_identity = deepcopy(snapshot.decision_rows[0])
    other_identity.sku = 'OTHER'
    snapshot.decision_rows += (other_identity,)
    store = AnalysisSnapshotStore()
    monkeypatch.setattr(api, 'ANALYSIS_STORE', store)
    store.put(snapshot)
    result = client.post('/api/economics/export', json=BODY)
    assert result.status_code == 400
    assert result.json()['error']['code'] == 'ECONOMICS_EXPORT_IDENTITY_CONFLICT'


def test_save_failure_keeps_existing_file_and_returns_retryable_error(tmp_path, monkeypatch):
    path = tmp_path / 'project.json'
    monkeypatch.setattr(api, 'PROJECT_PATH', path)
    assert client.put('/api/project/cost-prices/A', json={'cost': '40'}).status_code == 200
    previous = path.read_bytes()
    def fail(*_): raise OSError('test write failure')
    monkeypatch.setattr(api, 'save_project_atomic', fail)
    result = client.put('/api/project/cost-prices/A', json={'cost': '50'})
    assert result.status_code == 500
    assert result.json()['error']['code'] == 'COST_PRICE_SAVE_FAILED'
    assert path.read_bytes() == previous


def test_same_value_becoming_manual_invalidates_inflight_analysis(tmp_path, monkeypatch):
    from backend.cost_prices import cost_fingerprint
    from backend.shipment.api_context import ShipmentPreparationError
    path = tmp_path / 'project.json'
    monkeypatch.setattr(api, 'PROJECT_PATH', path)
    response = client.post('/api/analysis', files=_analysis_files(), data=_analysis_data())
    assert response.status_code == 200
    snapshot = api.ANALYSIS_STORE.latest()
    expected = cost_fingerprint(load_project(path))
    result = client.put('/api/project/cost-prices/ART-1', json={'cost': '100'})
    assert result.status_code == 200
    assert result.json()['changed'] is True
    with pytest.raises(ShipmentPreparationError, match='COST_PRICES_CHANGED_DURING_ANALYSIS'):
        api.commit_analysis_snapshot_if_current(snapshot,
            expected_pack_fingerprint=api.pack_multiplicity_fingerprint(load_project(path)),
            expected_cost_fingerprint=expected)


def test_cost_change_during_analysis_rejects_snapshot_commit(tmp_path, monkeypatch):
    path = tmp_path / 'project.json'
    monkeypatch.setattr(api, 'PROJECT_PATH', path)
    real_analyze = api.analyze
    def change_while_calculating(*args, **kwargs):
        result = real_analyze(*args, **kwargs)
        assert client.put('/api/project/cost-prices/ART-1', json={'cost': '150'}).status_code == 200
        return result
    monkeypatch.setattr(api, 'analyze', change_while_calculating)
    result = client.post('/api/analysis', files=_analysis_files(), data=_analysis_data())
    assert result.status_code == 409, result.text
    assert result.json()['error']['code'] == 'COST_PRICES_CHANGED_DURING_ANALYSIS'
    assert api.ANALYSIS_STORE.latest() is None
    assert load_project(path).cost_prices['ART-1'].cost == Decimal('150')


def test_excel_race_cannot_download_an_old_cost_scenario(tmp_path, monkeypatch):
    monkeypatch.setattr(api, 'PROJECT_PATH', tmp_path / 'project.json')
    store = AnalysisSnapshotStore()
    monkeypatch.setattr(api, 'ANALYSIS_STORE', store)
    snapshot = sample_snapshot()
    snapshot.snapshot_id = 'snap-1'
    store.put(snapshot)
    real_export = api.export_economics
    def change_during_export(report):
        data = real_export(report)
        assert client.put('/api/project/cost-prices/26572', json={'cost': '75'}).status_code == 200
        return data
    monkeypatch.setattr(api, 'export_economics', change_during_export)
    result = client.post('/api/economics/export', json=BODY)
    assert result.status_code == 409
    assert result.json()['error']['code'] == 'ECONOMICS_INPUT_CHANGED'
