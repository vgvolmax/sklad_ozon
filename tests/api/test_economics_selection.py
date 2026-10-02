"""One server-owned selection for the screen, summary and downloaded workbook."""
from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import backend.api as api
from backend.advertising_store import AdvertisingData, save_advertising
from backend.domain.advertising import AdvertisingDay, AdvertisingCampaign
from backend.main import app
from backend.shipment.store import AnalysisSnapshotStore
from tests.api.test_economics_workspace import BODY
from tests.economics.test_workspace import sample_snapshot

CLIENT = TestClient(app)


@pytest.fixture
def selected_snapshot(tmp_path, monkeypatch):
    monkeypatch.setattr(api, 'PROJECT_PATH', tmp_path / 'project.json')
    monkeypatch.setattr(api, 'ANALYSIS_STORE', AnalysisSnapshotStore())
    snap = sample_snapshot()
    snap.snapshot_id = 'snap-1'
    other = deepcopy(snap.unit_economics[0]); other.sku = 'OTHER'; other.cost = Decimal('20')
    identity = deepcopy(snap.decision_rows[0]); identity.sku = 'OTHER'; identity.product_name = 'Другой товар'
    route = deepcopy(snap.route_economics[0]); route.sku = 'OTHER'; route.observed_qty = 2
    snap.unit_economics += (other,)
    snap.decision_rows += (identity,)
    snap.route_economics += (route,)
    api.ANALYSIS_STORE.put(snap)
    return snap


def report(**extra):
    result = CLIENT.post('/api/economics/workspace', json={**BODY, **extra})
    assert result.status_code == 200, result.text
    return result.json()['workspace']


def test_lower_filter_and_export_have_the_same_sku_set(selected_snapshot):
    r = report(filter='lower')
    assert [p['sku'] for p in r['products']] == ['OTHER']
    assert r['catalog_product_count'] == 2
    response = CLIENT.post('/api/economics/export', json={**BODY, 'filter': 'lower'})
    assert response.status_code == 200, response.text
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.max_row == 2 and sheet.cell(2, 15).value == 'OTHER'


def test_search_matches_cluster_and_name_case_insensitively(selected_snapshot):
    assert [p['sku'] for p in report(search='ДРУГОЙ')['products']] == ['OTHER']
    assert len(report(search='казань')['products']) == 1


def test_empty_selection_has_no_download(selected_snapshot):
    assert report(search='нет такого')['products'] == []
    response = CLIENT.post('/api/economics/export', json={**BODY, 'search': 'нет такого'})
    assert response.status_code == 400
    assert response.json()['error']['code'] == 'ECONOMICS_EXPORT_EMPTY'


@pytest.mark.parametrize('extra', [{'filter': 'typo'}, {'filter': None},
    {'search': []}, {'search': 'x' * 201}])
def test_invalid_selection_is_rejected(selected_snapshot, extra):
    response = CLIENT.post('/api/economics/workspace', json={**BODY, **extra})
    assert response.status_code == 400


def test_duplicate_and_blank_articles_do_not_merge_or_block_sku_rows(selected_snapshot):
    selected_snapshot.decision_rows[1].article = ''
    response = CLIENT.post('/api/economics/export', json=BODY)
    assert response.status_code == 200, response.text
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.max_row == 3
    assert {sheet.cell(r, 15).value for r in (2, 3)} == {'SKU', 'OTHER'}
    selected_snapshot.decision_rows[1].article = '26572'
    response = CLIENT.post('/api/economics/export', json=BODY)
    assert response.status_code == 200
    assert load_workbook(BytesIO(response.content)).active.max_row == 3


def test_formula_looking_identity_stays_text(selected_snapshot):
    selected_snapshot.decision_rows[1].article = '=1+1'
    response = CLIENT.post('/api/economics/export', json=BODY)
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    cell = next(sheet.cell(r, 1) for r in (2, 3) if sheet.cell(r, 15).value == 'OTHER')
    assert cell.value == '=1+1' and cell.data_type == 's'


@pytest.mark.parametrize('spend,expected_after', [('0', '170'), ('50', '120')])
def test_summary_uses_known_spend_even_when_drr_is_unknown(selected_snapshot, spend, expected_after):
    save_advertising(api.PROJECT_PATH.with_name('advertising.json'), AdvertisingData(
        (AdvertisingDay('1', 'SKU', date(2026, 9, 20), Decimal(spend)),),
        (AdvertisingCampaign('1', 'test.xlsx', '2026-10-02T08:00:00+00:00'),)))
    r = report(search='SKU')
    assert r['products'][0]['real_drr_rate'] is None
    totals = r['totals']
    assert totals['profit_before_ads'] == '170'
    assert totals['advertising_spend'] == spend
    assert totals['profit_after_uploaded_ads'] == expected_after
    assert totals['qty'] == totals['covered_qty'] == 15


def test_missing_ads_are_unknown_and_partial_routes_only_contribute_known_profit(selected_snapshot):
    selected_snapshot.route_economics[1].current_profit_per_unit = None
    selected_snapshot.route_economics[1].reason_codes = ('MISSING_ROUTE',)
    t = report(search='SKU')['totals']
    assert t['profit_before_ads'] == '180'
    assert t['advertising_spend'] is None and t['profit_after_uploaded_ads'] is None
    assert t['covered_qty'] == 10 and t['qty'] == 15 and t['profit_partial']


def test_summary_clips_spend_to_the_same_selected_period(selected_snapshot):
    from tests.economics.test_period import period_snapshot
    selected_snapshot.economics_period_evidence = period_snapshot().economics_period_evidence
    save_advertising(api.PROJECT_PATH.with_name('advertising.json'), AdvertisingData(
        (AdvertisingDay('1', 'SKU', date(2026, 9, 1), Decimal('900')),
         AdvertisingDay('1', 'SKU', date(2026, 9, 2), Decimal('10'))),
        (AdvertisingCampaign('1', 'test.xlsx', '2026-10-02T08:00:00+00:00'),)))
    t = report(search='SKU', period_from='2026-09-02', period_to='2026-09-02')['totals']
    assert t['profit_before_ads'] == '-6'
    assert t['advertising_spend'] == '10' and t['profit_after_uploaded_ads'] == '-16'
    assert t['qty'] == t['covered_qty'] == 3
