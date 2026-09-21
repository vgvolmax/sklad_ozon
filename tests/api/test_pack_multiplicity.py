from io import BytesIO
from types import SimpleNamespace
from openpyxl import Workbook
from fastapi.testclient import TestClient
import pytest

import backend.api as api
from backend.main import app
from backend.ozon.contracts import OzonCredentialContext, OzonCredentials
from backend.project import PackMultiplicityRecord, Project, load_project, save_project_atomic
from backend.ingestion.supplier_packaging import PackMultiplicityEvidence
from backend.pack_multiplicity import pack_multiplicity_fingerprint, resolve_pack_multiplicity
from backend.shipment.api_context import ShipmentPreparationError
from tests.helpers.xlsx_fixtures import make_real_unitka

client=TestClient(app)

def xlsx(rows):
    book=Workbook(); sheet=book.active
    for row in rows: sheet.append(row)
    stream=BytesIO(); book.save(stream); return stream.getvalue()

def rtp_xlsx(rows):
    book=Workbook(); book.active.title='Обложка'
    sheet=book.create_sheet('Прайс списком')
    sheet.append(['Прайс']); sheet.append([]); sheet.append(['КОД','Упак'])
    for row in rows: sheet.append(row)
    stream=BytesIO(); book.save(stream); return stream.getvalue()

def test_crud_import_export_and_persistence(tmp_path,monkeypatch):
    path=tmp_path/'project.json'; monkeypatch.setattr(api,'PROJECT_PATH',path)
    save_project_atomic(path,Project(pack_multiplicity={'17261':PackMultiplicityRecord(20)}))
    assert client.get('/api/project/pack-multiplicity').json()['items'][0]['source']=='unitka'
    response=client.put('/api/project/pack-multiplicity/17261',json={'pack_multiple':50})
    assert response.status_code==200 and response.json()['item']['source']=='manual'
    assert load_project(path).pack_multiplicity['17261'].override_pack_multiple==50
    response=client.post('/api/project/pack-multiplicity/import',files={'file':('packs.xlsx',xlsx([['Артикул','Кратность'],['39439',8],['bad',0]]),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.json()['accepted']==1 and response.json()['rejected']==1
    assert load_project(path).pack_multiplicity['17261'].override_pack_multiple==50
    assert client.delete('/api/project/pack-multiplicity/17261').json()['item']['pack_multiple']==20
    exported=client.get('/api/project/pack-multiplicity/export')
    assert exported.status_code==200 and exported.content.startswith(b'PK')


def test_mutations_clear_derived_analysis_and_shipment_stores(tmp_path,monkeypatch):
    path=tmp_path/'project.json'; monkeypatch.setattr(api,'PROJECT_PATH',path)
    save_project_atomic(path,Project(pack_multiplicity={'17261':PackMultiplicityRecord(20)}))
    cleared=[]
    monkeypatch.setattr(api.ANALYSIS_STORE,'clear',lambda:cleared.append('analysis'))
    monkeypatch.setattr(api.SHIPMENT_PLAN_STORE,'clear',lambda:cleared.append('shipment'))
    assert client.put('/api/project/pack-multiplicity/17261',json={'pack_multiple':50}).status_code==200
    assert cleared == ['analysis','shipment']


def test_delete_and_xlsx_import_mutations_clear_derived_stores(tmp_path,monkeypatch):
    path=tmp_path/'project.json'; monkeypatch.setattr(api,'PROJECT_PATH',path)
    save_project_atomic(path,Project(pack_multiplicity={
        '17261':PackMultiplicityRecord(unitka_pack_multiple=20,override_pack_multiple=50,override_origin='manual',override_updated_at='now')}))
    cleared=[]
    monkeypatch.setattr(api.ANALYSIS_STORE,'clear',lambda:cleared.append('analysis'))
    monkeypatch.setattr(api.SHIPMENT_PLAN_STORE,'clear',lambda:cleared.append('shipment'))

    assert client.delete('/api/project/pack-multiplicity/17261').status_code==200
    assert cleared==['analysis','shipment']
    cleared.clear()
    response=client.post('/api/project/pack-multiplicity/import',files={
        'file':('packs.xlsx',xlsx([['Артикул','Кратность'],['39439',8]]))})
    assert response.status_code==200
    assert cleared==['analysis','shipment']


def test_real_rtp_price_import_precedence_snapshot_and_noop(tmp_path,monkeypatch):
    path=tmp_path/'project.json'; monkeypatch.setattr(api,'PROJECT_PATH',path)
    save_project_atomic(path,Project(pack_multiplicity={
        '28200':PackMultiplicityRecord(unitka_pack_multiple=10,
            override_pack_multiple=20,override_origin='manual',override_updated_at='now'),
        'OLD':PackMultiplicityRecord(rtp_price_pack_multiple=30,
            rtp_price_updated_at='old'),
        '51710':PackMultiplicityRecord(rtp_price_pack_multiple=100,
            rtp_price_updated_at='old'),
        'CONFLICT':PackMultiplicityRecord(rtp_price_pack_multiple=50,
            rtp_price_updated_at='old'),
    }))
    cleared=[]
    monkeypatch.setattr(api.ANALYSIS_STORE,'clear',lambda:cleared.append('analysis'))
    monkeypatch.setattr(api.SHIPMENT_PLAN_STORE,'clear',lambda:cleared.append('shipment'))
    price=rtp_xlsx([
        ['28200','15/1'],['28201','18/1'],['28202','14/1'],
        ['28206','14/1'],['29352','150/10'],['34886','70/1'],
        ['51710','100+/1'],['SAME','72/6'],['SAME','72/12'],
        ['CONFLICT','72/6'],['CONFLICT','24/12'],
    ])
    first=client.post('/api/project/pack-multiplicity/import',files={'file':('rtp.xlsx',price)}).json()
    assert (first['source_format'],first['accepted'],first['changed']) == ('rtp_price',7,True)
    assert first['rejected']==2
    project=load_project(path)
    assert project.pack_multiplicity['OLD'].rtp_price_pack_multiple is None
    assert project.pack_multiplicity['51710'].rtp_price_pack_multiple is None
    assert project.pack_multiplicity['CONFLICT'].rtp_price_pack_multiple is None
    assert {a:project.pack_multiplicity[a].rtp_price_pack_multiple for a in
            ('28200','28201','28202','28206','29352','34886','SAME')} == {
            '28200':15,'28201':18,'28202':14,'28206':14,
            '29352':150,'34886':70,'SAME':72}
    assert resolve_pack_multiplicity(project.pack_multiplicity['28200']).pack_multiple==20
    assert cleared==['analysis','shipment']
    stamp=project.pack_multiplicity['28200'].rtp_price_updated_at
    cleared.clear()
    second=client.post('/api/project/pack-multiplicity/import',files={'file':('rtp.xlsx',price)}).json()
    assert second['changed'] is False and cleared==[]
    assert load_project(path).pack_multiplicity['28200'].rtp_price_updated_at==stamp


def test_unitka_change_invalidates_but_noop_preserves_derived_stores(tmp_path,monkeypatch):
    path=tmp_path/'project.json'; monkeypatch.setattr(api,'PROJECT_PATH',path)
    project=Project(pack_multiplicity={'17261':PackMultiplicityRecord(20)})
    save_project_atomic(path,project)
    cleared=[]
    monkeypatch.setattr(api.ANALYSIS_STORE,'clear',lambda:cleared.append('analysis'))
    monkeypatch.setattr(api.SHIPMENT_PLAN_STORE,'clear',lambda:cleared.append('shipment'))

    with api.PROJECT_PERSISTENCE_LOCK:
        unchanged,changed=api._persist_unitka_baseline(project,(
            PackMultiplicityEvidence('17261',20,2,'x / 20',()),))
    assert changed is False and unchanged is project and cleared==[]

    with api.PROJECT_PERSISTENCE_LOCK:
        updated,changed=api._persist_unitka_baseline(project,(
            PackMultiplicityEvidence('17261',50,2,'x / 50',()),))
    assert changed is True and updated.pack_multiplicity['17261'].unitka_pack_multiple==50
    assert cleared==['analysis','shipment']


def test_unitka_import_endpoint_invalidates_only_for_baseline_change(tmp_path,monkeypatch):
    path=tmp_path/'project.json'; monkeypatch.setattr(api,'PROJECT_PATH',path)
    save_project_atomic(path,Project(pack_multiplicity={
        '17261':PackMultiplicityRecord(20)}))
    cleared=[]
    monkeypatch.setattr(api.ANALYSIS_STORE,'clear',lambda:cleared.append('analysis'))
    monkeypatch.setattr(api.SHIPMENT_PLAN_STORE,'clear',lambda:cleared.append('shipment'))

    same=make_real_unitka(pack_rows=[['17261','20/1']])
    assert client.post('/api/import/unitka',files={'file':('unitka.xlsx',same)}).status_code==200
    assert cleared==[]

    changed=make_real_unitka(pack_rows=[['17261','50/1']])
    assert client.post('/api/import/unitka',files={'file':('unitka.xlsx',changed)}).status_code==200
    assert cleared==['analysis','shipment']


def test_failed_unitka_persistence_does_not_invalidate_stores(tmp_path,monkeypatch):
    monkeypatch.setattr(api,'PROJECT_PATH',tmp_path/'project.json')
    project=Project(pack_multiplicity={'17261':PackMultiplicityRecord(20)})
    cleared=[]
    monkeypatch.setattr(api,'save_project_atomic',lambda *_: (_ for _ in ()).throw(OSError('disk')))
    monkeypatch.setattr(api.ANALYSIS_STORE,'clear',lambda:cleared.append('analysis'))
    monkeypatch.setattr(api.SHIPMENT_PLAN_STORE,'clear',lambda:cleared.append('shipment'))
    with pytest.raises(OSError,match='disk'):
        api._persist_unitka_baseline(project,(
            PackMultiplicityEvidence('17261',50,2,'x / 50',()),))
    assert cleared==[]


def test_stale_analysis_commit_is_rejected_without_store_write(tmp_path,monkeypatch):
    path=tmp_path/'project.json'; monkeypatch.setattr(api,'PROJECT_PATH',path)
    old=Project(pack_multiplicity={'17261':PackMultiplicityRecord(20)})
    save_project_atomic(path,old)
    expected=pack_multiplicity_fingerprint(old)
    save_project_atomic(path,Project(pack_multiplicity={
        '17261':PackMultiplicityRecord(50)}))
    writes=[]
    monkeypatch.setattr(api.ANALYSIS_STORE,'put',lambda snapshot:writes.append(snapshot))

    with pytest.raises(ShipmentPreparationError) as caught:
        api.commit_analysis_snapshot_if_current(
            object(),expected_pack_fingerprint=expected)

    assert caught.value.code=='PACK_MULTIPLICITY_CHANGED_DURING_ANALYSIS'
    assert writes==[]


def _shipment_commit_context():
    return OzonCredentialContext('credential-context',OzonCredentials('client','key'),1)


def _configure_shipment_commit(monkeypatch, snapshot):
    if snapshot is not None and snapshot.shippable_plan is not None:
        snapshot.shippable_plan.lines=()
    monkeypatch.setattr(api.OZON_VAULT,'is_context_active',lambda context:True)
    monkeypatch.setattr(api.ANALYSIS_STORE,'get',lambda snapshot_id:snapshot)
    monkeypatch.setattr(api.ANALYSIS_STORE,'latest',
                        lambda:SimpleNamespace(snapshot_id='A1'))
    original = api.require_current_working_plan
    def require(**kwargs):
        snap = api._require_current_shippable_plan(kwargs['analysis_snapshot_id'], kwargs['shippable_plan_id'])
        return snap, snap.shippable_plan, SimpleNamespace(working_plan_id='WP1')
    monkeypatch.setattr(api, 'require_current_working_plan', require)


def test_stale_shipment_commit_rejects_changed_pack_before_store_write(tmp_path,monkeypatch):
    path=tmp_path/'project.json'; monkeypatch.setattr(api,'PROJECT_PATH',path)
    old=Project(pack_multiplicity={'17261':PackMultiplicityRecord(20)})
    save_project_atomic(path,old)
    expected=pack_multiplicity_fingerprint(old)
    snapshot=SimpleNamespace(shippable_plan=SimpleNamespace(shippable_plan_id='P1'))
    _configure_shipment_commit(monkeypatch,snapshot)
    save_project_atomic(path,Project(pack_multiplicity={'17261':PackMultiplicityRecord(50)}))
    writes=[]; monkeypatch.setattr(api.SHIPMENT_PLAN_STORE,'put',writes.append)

    with pytest.raises(ShipmentPreparationError) as caught:
        api.commit_shipment_plan_if_current(
            object(),expected_analysis_snapshot_id='A1',
            expected_shippable_plan_id='P1',expected_pack_fingerprint=expected,
            expected_working_plan_id='WP1',
            credential_context=_shipment_commit_context())

    assert caught.value.code=='PACK_MULTIPLICITY_CHANGED_DURING_SHIPMENT_VALIDATION'
    assert writes==[]


@pytest.mark.parametrize('snapshot',[None,SimpleNamespace(shippable_plan=None),
    SimpleNamespace(shippable_plan=SimpleNamespace(shippable_plan_id='P2'))])
def test_stale_shipment_commit_rejects_changed_analysis_snapshot(tmp_path,monkeypatch,snapshot):
    path=tmp_path/'project.json'; monkeypatch.setattr(api,'PROJECT_PATH',path)
    project=Project(pack_multiplicity={'17261':PackMultiplicityRecord(20)})
    save_project_atomic(path,project)
    _configure_shipment_commit(monkeypatch,snapshot)
    writes=[]; monkeypatch.setattr(api.SHIPMENT_PLAN_STORE,'put',writes.append)

    with pytest.raises(ShipmentPreparationError) as caught:
        api.commit_shipment_plan_if_current(
            object(),expected_analysis_snapshot_id='A1',
            expected_shippable_plan_id='P1',
            expected_pack_fingerprint=pack_multiplicity_fingerprint(project),
            expected_working_plan_id='WP1',
            credential_context=_shipment_commit_context())

    assert caught.value.code=='SHIPMENT_INPUT_CHANGED'
    assert writes==[]


def test_current_shipment_commit_writes_once(tmp_path,monkeypatch):
    path=tmp_path/'project.json'; monkeypatch.setattr(api,'PROJECT_PATH',path)
    project=Project(pack_multiplicity={'17261':PackMultiplicityRecord(20)})
    save_project_atomic(path,project)
    _configure_shipment_commit(monkeypatch,
        SimpleNamespace(shippable_plan=SimpleNamespace(shippable_plan_id='P1')))
    shipment=object(); writes=[]; monkeypatch.setattr(api.SHIPMENT_PLAN_STORE,'put',writes.append)

    api.commit_shipment_plan_if_current(
        shipment,expected_analysis_snapshot_id='A1',expected_shippable_plan_id='P1',
        expected_pack_fingerprint=pack_multiplicity_fingerprint(project),
        expected_working_plan_id='WP1',
            credential_context=_shipment_commit_context())

    assert writes==[shipment]


def test_stale_shipment_commit_preserves_credential_context_guard(tmp_path,monkeypatch):
    path=tmp_path/'project.json'; monkeypatch.setattr(api,'PROJECT_PATH',path)
    project=Project(pack_multiplicity={'17261':PackMultiplicityRecord(20)})
    save_project_atomic(path,project)
    monkeypatch.setattr(api.OZON_VAULT,'is_context_active',lambda context:False)
    writes=[]; monkeypatch.setattr(api.SHIPMENT_PLAN_STORE,'put',writes.append)

    with pytest.raises(ShipmentPreparationError) as caught:
        api.commit_shipment_plan_if_current(
            object(),expected_analysis_snapshot_id='A1',
            expected_shippable_plan_id='P1',
            expected_pack_fingerprint=pack_multiplicity_fingerprint(project),
            expected_working_plan_id='WP1',
            credential_context=_shipment_commit_context())

    assert caught.value.code=='OZON_CREDENTIAL_CONTEXT_CHANGED'
    assert writes==[]
