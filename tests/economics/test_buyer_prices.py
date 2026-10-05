from dataclasses import replace
from datetime import date
from decimal import Decimal

import pytest

from backend.analytics._weeks import ObservationCoverage
from backend.economics.daily_series import build_daily_evidence
from backend.economics.export import export_economics
from backend.economics.workspace import build_economics_workspace
from backend.project import EconomicsSettings
from tests.economics.test_daily_series import order
from tests.economics.test_period import period_snapshot
from tests.economics.test_workspace import sample_snapshot


def workspace(snap, **kwargs):
    kwargs.setdefault('planned_drr', '0')
    return build_economics_workspace(snap, margin='.20', roi='.40', goal='margin', **kwargs)


def priced_snapshot(orders, *, complete=True):
    snap = period_snapshot()
    d = Decimal
    snap.economics_settings = EconomicsSettings(d(0), d(0), d(1), d(0),
        'usn_income', d(0), d(0), d(0))
    snap.unit_economics[0].cost = d(80)
    snap.unit_economics[0].commission = d(0)
    for route in snap.route_economics:
        route.route_cost_rub = d(0)
    snap.daily_order_evidence = build_daily_evidence(orders,
        ObservationCoverage(date(2026, 9, 1), date(2026, 9, 30)),
        orders_complete=complete)
    return snap


def prices(snap, **kwargs):
    return workspace(snap, planned_drr='0', **kwargs)['products'][0]['buyer_prices']


def test_period_prices_weight_units_and_each_orders_spp_not_daily_means_or_price_ratio():
    snap = priced_snapshot([order('2026-09-01', 100, 40, 2, 'SKU'),
        order('2026-09-01', 200, 100, 1, 'SKU'),
        order('2026-09-02', 100, 50, 1, 'SKU')])
    p = prices(snap)
    assert p['buyer_price_mean'] == Decimal('57.5')
    assert p['spp_mean'] == Decimal('.55')
    assert p['target_buyer_price'] == Decimal('45.00')
    assert p['ordered_qty'] == p['buyer_priced_qty'] == p['spp_priced_qty'] == 4
    assert p['complete'] is True
    selected = prices(snap, period_from='2026-09-02', period_to='2026-09-02')
    assert selected['buyer_price_mean'] == 50
    assert selected['spp_mean'] == Decimal('.5')
    assert selected['target_buyer_price'] == Decimal('50.00')
    assert selected['ordered_qty'] == 1


@pytest.mark.parametrize('paid,spp,target', [(100, '0', '100.00'), (0, '1', '0.00')])
def test_no_discount_and_free_buyer_price_are_known_values(paid, spp, target):
    p = prices(priced_snapshot([order('2026-09-01', 100, paid, 1, 'SKU')]))
    assert p['buyer_price_mean'] == paid
    assert p['spp_mean'] == Decimal(spp)
    assert p['target_buyer_price'] == Decimal(target)


def test_client_projection_rounds_half_kopek_up_without_rounding_spp_first():
    p = prices(priced_snapshot([order('2026-09-01', 100000, 105, 1, 'SKU')]))
    assert p['spp_mean'] == Decimal('.99895')
    assert p['target_buyer_price'] == Decimal('.11')


def test_buyer_price_and_spp_have_independent_coverage_and_history_completeness():
    snap = priced_snapshot([order('2026-09-01', 100, 40, 2, 'SKU'),
        order('2026-09-01', 0, 100, 1, 'SKU'),
        order('2026-09-02', 100, None, 3, 'SKU')], complete=False)
    p = prices(snap)
    assert p['buyer_price_mean'] == 60 and p['spp_mean'] == Decimal('.6')
    assert p['ordered_qty'] == 6 and p['buyer_priced_qty'] == 3 and p['spp_priced_qty'] == 2
    assert p['complete'] is False
    assert p['target_buyer_price'] == Decimal('40.00')
    snap.daily_order_evidence = build_daily_evidence([order('2026-09-01', 0, 40, 1, 'SKU')],
        ObservationCoverage(date(2026, 9, 1), date(2026, 9, 30)))
    p = prices(snap)
    assert p['buyer_price_mean'] == 40 and p['spp_mean'] is None
    assert p['target_buyer_price'] is None


def test_missing_history_empty_period_and_unreachable_target_do_not_invent_prices():
    assert prices(sample_snapshot())['buyer_price_mean'] is None
    snap = priced_snapshot([order('2026-09-01', 100, 40, 1, 'SKU')])
    empty = prices(snap, period_from='2026-09-03', period_to='2026-09-29')
    assert empty['ordered_qty'] == 0 and empty['target_buyer_price'] is None
    snap.route_economics[0].current_profit_per_unit = None
    snap.route_economics[0].reason_codes = ('MISSING_ROUTE',)
    p = prices(snap)
    assert p['buyer_price_mean'] == 40 and p['spp_mean'] == Decimal('.6')
    assert p['target_buyer_price'] is None


def test_same_article_skus_keep_distinct_prices_and_snapshot_is_immutable():
    snap = priced_snapshot([order('2026-09-01', 100, 40, 1, 'SKU'),
        order('2026-09-01', 100, 90, 1, 'OTHER')])
    from types import SimpleNamespace as NS
    snap.decision_rows += (NS(sku='OTHER', article='26572', product_name='Другой SKU'),)
    original = snap.daily_order_evidence
    report = workspace(snap, planned_drr='0')
    by_sku = {p['sku']: p for p in report['products']}
    assert by_sku['SKU']['buyer_prices']['buyer_price_mean'] == 40
    assert by_sku['OTHER']['buyer_prices']['buyer_price_mean'] == 90
    assert by_sku['OTHER']['buyer_prices']['target_buyer_price'] is None
    assert snap.daily_order_evidence is original


def test_current_enriched_evidence_replaces_missing_original_prices_without_mutation():
    snap = priced_snapshot([order('2026-09-01', 100, None, 1, 'SKU')])
    fresh = replace(snap.daily_order_evidence, days=build_daily_evidence(
        [order('2026-09-01', 100, 60, 1, 'SKU')],
        ObservationCoverage(date(2026, 9, 1), date(2026, 9, 30))).days)
    p = workspace(snap, planned_drr='0', daily_evidence=fresh)['products'][0]['buyer_prices']
    assert p['buyer_price_mean'] == 60 and p['target_buyer_price'] == Decimal('60.00')
    assert snap.daily_order_evidence.days[0].buyer_price_mean is None


def test_xlsx_keeps_numeric_prices_spp_and_coverage_for_selected_period():
    from io import BytesIO
    from openpyxl import load_workbook
    snap = priced_snapshot([order('2026-09-01', 100, 40, 2, 'SKU'),
        order('2026-09-02', 100, 80, 3, 'SKU')])
    report = workspace(snap, planned_drr='0', period_from='2026-09-02', period_to='2026-09-02')
    sheet = load_workbook(BytesIO(export_economics(report))).active
    cells = {h.value: sheet.cell(2, h.column) for h in sheet[1]}
    assert cells['Средняя цена клиента, ₽'].value == 80
    assert cells['Средний СПП, %'].value == .2
    assert cells['Средний СПП, %'].number_format == '0.0%'
    assert cells['Цена клиента при цели, ₽'].value == 80
    assert cells['Заказано для средних, шт.'].value == 3
    assert cells['Цена клиента известна, шт.'].value == 3
    assert cells['СПП известен, шт.'].value == 3
