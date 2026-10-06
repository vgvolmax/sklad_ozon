from datetime import date
from decimal import Decimal as D
from types import SimpleNamespace as NS

import pytest
import backend.application
from backend.economics.buyouts import FinanceSnapshot, FinanceProductLine, FinanceExpense

PERIOD = {'from': date(2026, 9, 1), 'to': date(2026, 9, 30)}


def basis():
    return [dict(sku='A', article='A', name='One', price=D('1000'),
                 profit_per_unit_before_ads=D('150'), pricing_complete=True),
            dict(sku='B', article='B', name='Two', price=D('500'),
                 profit_per_unit_before_ads=D('80'), pricing_complete=True)]


def finance(lines=None, expenses=None):
    return FinanceSnapshot('F', 'account', PERIOD['from'], PERIOD['to'],
        tuple(lines if lines is not None else [
            FinanceProductLine(date(2026, 9, 1), 'A', 70, D('3'), D('1')),
            FinanceProductLine(date(2026, 9, 2), 'B', 30, D('4'), D('2'))]),
        tuple(expenses if expenses is not None else [
            FinanceExpense(date(2026, 9, 3), 'advertising', 'Ads A', D('3000'), 'A'),
            FinanceExpense(date(2026, 9, 3), 'advertising', 'Ads without sales', D('2000'), 'NO-SALES'),
            FinanceExpense(date(2026, 9, 3), 'storage', 'Storage', D('2000'), 'B'),
            FinanceExpense(date(2026, 9, 3), 'commission', 'Commission', D('4000'))]), 'now')


def report(mode='orders', **kwargs):
    from backend.economics.period_profit import build_period_profit
    return build_period_profit(kwargs.pop('basis', basis()), PERIOD,
        [NS(sku='A', day=date(2026, 9, 1), quantity=100),
         NS(sku='B', day=date(2026, 9, 1), quantity=50)], kwargs.pop('finance', finance()),
        mode, **kwargs)


def test_only_quantity_changes_between_modes_and_no_sku_or_payout_ad_allocation():
    orders, buyouts = report(), report('buyouts')
    assert orders['totals']['profit_before_common'] == D('19000')
    assert orders['totals']['profit_after_common'] == D('12000')
    assert buyouts['totals']['profit_before_common'] == D('12900')
    assert buyouts['totals']['profit_after_common'] == D('5900')
    for r in (orders, buyouts):
        assert r['totals']['advertising_total'] == D('5000')
        assert r['totals']['other_common_total'] == D('2000')
    assert orders['totals']['margin'] == D('.096')


def test_filter_does_not_deduct_whole_shop_expenses_from_one_sku():
    r = report(selected_skus={'A'})
    assert [p['sku'] for p in r['products']] == ['A']
    assert r['totals']['profit_after_common'] == D('12000')
    assert r['selected_profit_before_common'] == D('15000')


def test_negative_net_return_is_preserved_and_margin_not_manufactured():
    f = finance([FinanceProductLine(date(2026, 9, 1), 'A', -2, D('-1'), D('-3'))], [])
    r = report('buyouts', finance=f)
    assert r['totals']['qty'] == -2
    assert r['totals']['profit_after_common'] == D('-300')
    assert r['totals']['margin'] is None


def test_unknown_return_quantity_and_old_sku_are_reported_as_partial():
    f = finance([FinanceProductLine(date(2026, 9, 1), 'OLD', None, None, D('-5'))], [])
    r = report('buyouts', finance=f)
    assert r['totals']['partial'] is True
    assert next(p for p in r['products'] if p['sku'] == 'OLD')['quantity_known'] is False
    assert r['totals']['profit_after_common'] is None


def test_missing_expenses_differs_from_complete_zero():
    assert report(finance=None)['totals']['profit_after_common'] is None
    assert report(finance=finance(expenses=[]))['totals']['profit_after_common'] == D('19000')


def test_missing_finance_does_not_mean_missing_unit_economics():
    r = report('buyouts', finance=None)
    assert r['totals']['profit_before_common'] is None
    assert r['coverage']['unit_available_count'] == 2
    assert r['coverage']['missing_unit_products'] == []
    assert [p['sku'] for p in r['coverage']['missing_quantity_products']] == ['A', 'B']
    assert all(p['unit_complete'] for p in r['products'])


def test_store_coverage_separates_quantity_from_unit_inputs_outside_filter():
    b = basis(); b[1].update(profit_per_unit_before_ads=None, pricing_complete=False, article='ART-B')
    f = finance([FinanceProductLine(date(2026, 9, 1), 'A', None, None, D('1')),
                 FinanceProductLine(date(2026, 9, 1), 'B', 3, D('1500'), D('2'))], [])
    r = report('buyouts', basis=b, finance=f, selected_skus={'A'})
    assert [p['sku'] for p in r['products']] == ['A']
    assert r['coverage']['missing_unit_products'] == [dict(sku='B', article='ART-B', name='Two')]
    assert r['coverage']['missing_quantity_products'] == [dict(sku='A', article='A', name='One')]
    assert r['coverage']['unit_available_count'] == 1
    assert r['totals']['profit_before_common'] is None


def test_known_zero_quantity_needs_no_unit_but_does_not_claim_a_unit_exists():
    b = basis(); b[1].update(profit_per_unit_before_ads=None, pricing_complete=False)
    r = report('buyouts', basis=b, finance=finance(lines=[], expenses=[]))
    assert r['totals']['profit_after_common'] == 0
    assert not r['totals']['partial']
    assert r['coverage']['missing_unit_products'] == []
    assert r['coverage']['missing_quantity_products'] == []
    assert r['coverage']['unit_available_count'] == 1


@pytest.mark.parametrize('shared_article', [False, True])
def test_export_distinguishes_unit_from_quantity_coverage_using_articles(shared_article):
    from io import BytesIO
    from openpyxl import load_workbook
    from backend.economics.export import export_period_profit
    b = basis(); b[0]['article'] = 'ART-A'; b[1].update(
        profit_per_unit_before_ads=None, pricing_complete=False, article='ART-B')
    if shared_article:
        b[0]['article'] = b[1]['article'] = 'SHARED'
    f = finance([FinanceProductLine(date(2026, 9, 1), 'A', None, None, D('1')),
                 FinanceProductLine(date(2026, 9, 1), 'B', 3, D('1500'), D('2'))], [])
    book = load_workbook(BytesIO(export_period_profit(report('buyouts', basis=b, finance=f))))
    summary = dict(book['Итог'].values)
    assert summary['Товары без полной юнитки'] == ('SHARED · SKU B' if shared_article else 'ART-B')
    assert summary['Товары с неизвестным количеством'] == ('SHARED · SKU A' if shared_article else 'ART-A')
    assert book['Товары']['I2'].value == 'Количество неизвестно'
    assert book['Товары']['I3'].value == 'Юнитка неполная'


def test_covered_duplicate_article_still_disambiguates_filtered_out_missing_unit():
    from io import BytesIO
    from openpyxl import load_workbook
    from backend.economics.export import export_period_profit
    b = basis(); b[0].update(profit_per_unit_before_ads=None, pricing_complete=False)
    b[0]['article'] = b[1]['article'] = 'SHARED'
    r = report(basis=b, selected_skus={'B'})
    assert r['coverage']['ambiguous_articles'] == ['SHARED']
    summary = dict(load_workbook(BytesIO(export_period_profit(r)))['Итог'].values)
    assert summary['Товары без полной юнитки'] == 'SHARED · SKU A'


def test_unknown_fee_and_negative_expense_credit_preserve_partial_provenance():
    f = finance(expenses=[FinanceExpense(date(2026, 9, 3), 'storage', 'Credit', D('-200')),
                         FinanceExpense(date(2026, 9, 3), 'other', 'Unknown', D('500'))])
    r = report(finance=f)
    assert r['totals']['profit_after_common'] == D('19200')
    assert r['totals']['partial'] is True
    assert any(e['role'] == 'unclassified' for e in r['expenses'])


def test_no_quantities_still_subtracts_known_expenses():
    r = report('buyouts', finance=finance(lines=[]))
    assert r['totals']['profit_before_common'] == 0
    assert r['totals']['profit_after_common'] == D('-7000')


def test_incomplete_unit_does_not_become_zero_profit():
    b = basis(); b[0].update(profit_per_unit_before_ads=None, pricing_complete=False)
    r = report(basis=b)
    assert r['totals']['partial'] is True
    assert r['totals']['profit_before_common'] == D('4000')
    assert r['totals']['profit_after_common'] == D('-3000')


def test_net_order_count_includes_in_progress_without_origin_and_excludes_cancelled():
    from backend.analytics.daily import build_daily_order_facts
    from backend.domain.contracts import OrderRecord, OrderLifecycle
    from backend.economics.period_profit import build_period_profit
    orders = tuple(OrderRecord('A',qty,origin,'B',lifecycle,'2026-09-01T10:00:00')
        for qty,origin,lifecycle in [(1,'A',OrderLifecycle.FULFILLED),
            (2,'',OrderLifecycle.IN_PROGRESS),(3,'A',OrderLifecycle.CANCELLED)])
    facts = build_daily_order_facts(orders,PERIOD['to'])
    r = build_period_profit(basis(),PERIOD,facts.demand.cells,finance(expenses=[]),'orders')
    assert r['totals']['qty'] == 3
    assert r['totals']['profit_before_common'] == D('450')
    assert sum(c.quantity for c in facts.fulfillment.cells) == 1
