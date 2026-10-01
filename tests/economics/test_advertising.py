from datetime import date
from decimal import Decimal
from io import BytesIO
import zipfile

from openpyxl import Workbook, load_workbook
import pytest

from backend.analytics._weeks import ObservationCoverage
from backend.domain.contracts import OrderLifecycle, OrderRecord, AnalysisSourceCoverage


def advertising_xlsx(*, campaign='123', start='01.09.2026', end='02.09.2026',
                     rows=None, dimension_broken=False, total=None):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Statistics'
    sheet.append([f'Период: {start} - {end}'])
    sheet.append(['Группировка: по дням'])
    sheet.append([f'Кампания по продвижению товаров № {campaign}. Отчет по товарам'])
    sheet.append(['День', 'SKU', 'Название товара', 'Расход, ₽, с НДС',
                  'Продажи в продвижении, ₽', 'ДРР в продвижении, %', 'ДРР, %'])
    rows = rows if rows is not None else [
        ['01.09.2026', 100, 'Кран', 10, 10, 1, .9],
        ['02.09.2026', 100, 'Кран', 30, 15, 2, .8],
    ]
    for row in rows:
        sheet.append(row)
    sheet.append(['Всего', None, None, total if total is not None else sum(r[3] for r in rows)])
    stream = BytesIO()
    book.save(stream)
    if not dimension_broken:
        return stream.getvalue()
    output = BytesIO()
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(output, 'w') as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                import re
                content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', content)
            target.writestr(item, content)
    return output.getvalue()


def test_daily_report_recovers_dimensions_and_uses_spend_not_attributed_rates():
    from backend.ingestion.advertising import import_advertising
    report = import_advertising(advertising_xlsx(dimension_broken=True), 'campaign.xlsx')
    assert report.campaign_id == '123'
    assert [(r.sku, r.day, r.spend) for r in report.days] == [
        ('100', date(2026, 9, 1), Decimal('10')),
        ('100', date(2026, 9, 2), Decimal('30')),
    ]


@pytest.mark.parametrize('rows,total', [
    ([['03.09.2026', 100, 'Кран', 10]], None),
    ([['01.09.2026', 100, 'Кран', -10]], None),
    ([['01.09.2026', True, 'Кран', 10]], None),
    ([['01.09.2026', 100, 'Кран', 10], ['01.09.2026', 100, 'Кран', 10]], None),
    ([['01.09.2026', 100, 'Кран', 10]], 20),
])
def test_invalid_advertising_never_silently_drops_money(rows, total):
    from backend.ingestion.advertising import import_advertising
    with pytest.raises(ValueError):
        import_advertising(advertising_xlsx(rows=rows, total=total), 'bad.xlsx')


def test_overlap_replaces_same_campaign_but_other_campaign_adds_expense(tmp_path):
    from backend.ingestion.advertising import import_advertising
    from backend.advertising_store import AdvertisingData, apply_report, load_advertising, save_advertising
    first = import_advertising(advertising_xlsx(), 'first.xlsx')
    data, result = apply_report(AdvertisingData(), first, {'100'})
    assert result['status'] == 'imported'
    same, result = apply_report(data, first, {'100'})
    assert same == data and result['status'] == 'duplicate'
    corrected = import_advertising(advertising_xlsx(start='02.09.2026', rows=[
        ['02.09.2026', 100, 'Кран', 20]]), 'corrected.xlsx')
    data, _ = apply_report(data, corrected, {'100'})
    another = import_advertising(advertising_xlsx(campaign='456', rows=[
        ['01.09.2026', 100, 'Кран', 5], ['02.09.2026', 100, 'Кран', 15]]), 'another.xlsx')
    data, _ = apply_report(data, another, {'100'})
    assert sum(r.spend for r in data.days) == Decimal('50')
    path = tmp_path / 'advertising.json'
    save_advertising(path, data)
    assert load_advertising(path) == data


def order(day, price, *, quantity=1, lifecycle=OrderLifecycle.FULFILLED, sku='100', channel='fbo'):
    return OrderRecord(sku, quantity, 'Москва', 'Москва', lifecycle,
                       accepted_at=day, seller_price=price, source_channel=channel)


def evidence(orders, **kwargs):
    from backend.economics.advertising import build_order_revenue
    return build_order_revenue(orders, ObservationCoverage(date(2026, 9, 1), date(2026, 9, 2)), **kwargs)


def test_total_drr_uses_all_orders_once_across_campaigns_and_not_a_mean():
    from backend.ingestion.advertising import import_advertising
    from backend.advertising_store import AdvertisingData, apply_report
    from backend.economics.advertising import real_drr_by_sku
    data, _ = apply_report(AdvertisingData(), import_advertising(advertising_xlsx(), 'a.xlsx'), {'100'})
    data, _ = apply_report(data, import_advertising(advertising_xlsx(campaign='456'), 'b.xlsx'), {'100'})
    # 80 advertising / (100 delivered + 200 in progress + 100 cancelled + 400 FBS) = 10%.
    source = evidence([
        order('2026-09-01', 100),
        order('2026-09-01', 200, lifecycle=OrderLifecycle.IN_PROGRESS),
        order('2026-09-01', 100, lifecycle=OrderLifecycle.CANCELLED),
        order('2026-09-02', 200, quantity=2, channel='fbs'),
        order('2026-08-31', 9999),
    ])
    product = real_drr_by_sku(data, source, {'100', '200'})['100']
    assert product['rate'] == Decimal('.1')
    assert product['spend'] == 80 and product['order_revenue'] == 800
    assert product['day_count'] == 2
    assert real_drr_by_sku(data, source, {'200'})['200']['rate'] is None


@pytest.mark.parametrize('case', ['missing_price', 'missing_date', 'incomplete_source', 'no_orders', 'outside_period'])
def test_unknown_revenue_or_period_never_becomes_zero_drr(case):
    from dataclasses import replace
    from backend.ingestion.advertising import import_advertising
    from backend.advertising_store import AdvertisingData, apply_report
    from backend.economics.advertising import real_drr_by_sku
    data, _ = apply_report(AdvertisingData(), import_advertising(advertising_xlsx(), 'a.xlsx'), {'100'})
    orders = [order('2026-09-01', 100)]
    kwargs = {}
    if case == 'missing_price': orders.append(order('2026-09-02', 0))
    if case == 'missing_date': orders.append(order('', 100))
    if case == 'incomplete_source': kwargs['source_coverage'] = AnalysisSourceCoverage(True, False, True, True)
    if case == 'no_orders': orders = []
    source = evidence(orders, **kwargs)
    if case == 'outside_period': source = replace(source, period_end=date(2026, 9, 1))
    result = real_drr_by_sku(data, source, {'100'})['100']
    assert result['rate'] is None and result['reason']


def test_zero_spend_is_known_and_real_drr_over_100_percent_is_not_clamped():
    from backend.ingestion.advertising import import_advertising
    from backend.advertising_store import AdvertisingData, apply_report
    from backend.economics.advertising import real_drr_by_sku
    for spend, expected in [(0, Decimal('0')), (300, Decimal('3'))]:
        report = import_advertising(advertising_xlsx(rows=[['01.09.2026', 100, 'Кран', spend]]), 'a.xlsx')
        data, _ = apply_report(AdvertisingData(), report, {'100'})
        assert real_drr_by_sku(data, evidence([order('2026-09-01', 100)]), {'100'})['100']['rate'] == expected


def test_vendor_rounding_drift_warns_and_keeps_per_day_costs():
    from backend.ingestion.advertising import import_advertising
    rows = [['01.09.2026', str(100 + i), 'Товар', 10] for i in range(20)]
    report = import_advertising(advertising_xlsx(rows=rows, total=199.9), 'rounded.xlsx')
    assert sum(day.spend for day in report.days) == Decimal('200')
    assert '0.10 ₽' in report.warning


def test_sparse_sheet_cell_limit_stops_before_materializing_entire_workbook(monkeypatch):
    from backend.ingestion.advertising import import_advertising
    import backend.ingestion.advertising as ingestion
    monkeypatch.setattr(ingestion, 'MAX_REPORT_CELLS', 20)
    with pytest.raises(ValueError, match='слишком много ячеек'):
        import_advertising(advertising_xlsx(), 'large.xlsx')


def test_corrupt_saved_money_has_actionable_error_and_is_not_overwritten(tmp_path):
    import json
    from backend.ingestion.advertising import import_advertising
    from backend.advertising_store import AdvertisingData, apply_report, save_advertising, load_advertising
    path = tmp_path / 'advertising.json'
    data, _ = apply_report(AdvertisingData(), import_advertising(advertising_xlsx(), 'a.xlsx'), {'100'})
    save_advertising(path, data)
    payload = json.loads(path.read_text())
    payload['days'][0]['spend'] = 'broken'
    path.write_text(json.dumps(payload))
    with pytest.raises(ValueError, match='сохранённые расходы'):
        load_advertising(path)


def test_disjoint_periods_do_not_add_revenue_from_the_gap():
    from backend.domain.advertising import AdvertisingDay
    from backend.advertising_store import AdvertisingData
    from backend.economics.advertising import real_drr_by_sku
    from backend.economics.advertising import build_order_revenue
    source = build_order_revenue([order('2026-09-01', 100), order('2026-09-02', 900), order('2026-09-03', 100)],
        ObservationCoverage(date(2026,9,1),date(2026,9,3)))
    costs = AdvertisingData(days=(AdvertisingDay('1', '100', date(2026,9,1),Decimal(10)),
                                 AdvertisingDay('2', '100', date(2026,9,3),Decimal(30))))
    actual = real_drr_by_sku(costs, source, {'100'})['100']
    assert actual['rate'] == Decimal('.2')
    assert actual['order_revenue'] == 200
    assert len(actual['periods']) == 2


@pytest.mark.parametrize('campaign', ['0', '1' * 33])
def test_campaign_id_is_validated_before_it_can_poison_persistence(campaign):
    from backend.ingestion.advertising import import_advertising
    with pytest.raises(ValueError, match='кампании'):
        import_advertising(advertising_xlsx(campaign=campaign), 'invalid.xlsx')


def test_report_without_total_cannot_certify_a_complete_period():
    from backend.ingestion.advertising import import_advertising
    book = load_workbook(BytesIO(advertising_xlsx()))
    book.active.delete_rows(book.active.max_row)
    buffer = BytesIO(); book.save(buffer)
    with pytest.raises(ValueError, match='Всего'):
        import_advertising(buffer.getvalue(), 'partial.xlsx')
